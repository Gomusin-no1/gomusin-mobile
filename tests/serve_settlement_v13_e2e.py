"""Disposable, synthetic-only local browser test server. Never loads a production DB."""
import os
import sys
from datetime import datetime
from pathlib import Path

os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
os.environ['ADMIN_USERNAME'] = ''
os.environ['ADMIN_PASSWORD'] = ''
os.environ['COOKIE_SECURE'] = 'false'
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_settlement_v13 import sales_book, gyeonggi_book, juan_book
from app import app, db, Branch, User
from werkzeug.security import generate_password_hash
import openpyxl
import io


if __name__ == '__main__':
    dest = Path(sys.argv[1]); dest.mkdir(parents=True, exist_ok=True)
    app.config.update(TESTING=True, SECRET_KEY='synthetic-browser-test-only')
    with app.app_context():
        db.create_all()
        branches=[Branch(name=n,code=f'TEST{i}',company_code='v13-test') for i,n in enumerate(['부평본점','간석점','숭의점'])]
        db.session.add_all(branches);db.session.flush()
        db.session.add(User(username='test-admin',display_name='검증관리자',password_hash=generate_password_hash('test-only-password'),role='admin',company_code='v13-test',branch_id=branches[0].id,email='test@example.invalid',email_verified_at=datetime.utcnow(),active=True))
        db.session.commit()
    rows=[]
    for i,(name,ph,extra) in enumerate([('테스트일치','01011112222',-30000),('테스트초과','01011113333',-50000),('테스트부족','01011114444',-20000)]):
        wb=openpyxl.load_workbook(io.BytesIO(sales_book(name=name,phone=ph)))
        if i==0:
            # Include KT and a missing settlement row in the same sales workbook.
            ws=wb.active
            for target,src in [(5,sales_book(vendor='주안')),(6,sales_book(name='테스트누락',phone='01011116666'))]:
                row=openpyxl.load_workbook(io.BytesIO(src)).active
                for c in range(1,32):ws.cell(target,c,row.cell(4,c).value)
        wb.save(dest/f'sales-{i}.xlsx');wb.close()
        v=openpyxl.load_workbook(io.BytesIO(gyeonggi_book(name=name,phone=ph,extra=extra)))
        rows.append([c.value for c in v.active[8]]);v.close()
    wb=openpyxl.load_workbook(io.BytesIO(gyeonggi_book()))
    ws=wb.active;ws.delete_rows(8,ws.max_row)
    for row in rows:ws.append(row)
    wb.save(dest/'gyeonggi.xlsx');wb.close()
    (dest/'juan.xlsx').write_bytes(juan_book())
    print('SYNTHETIC_E2E_READY',flush=True)
    app.run(host='127.0.0.1',port=5013,use_reloader=False)
