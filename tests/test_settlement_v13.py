"""Synthetic integration fixtures only. Customer source workbooks never enter git."""
import hashlib
import io
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl
import test_signature_and_tenant as fixtures
from app import app, db, User, AuditLog, SettlementV13Run, SettlementV13Source
from settlement_v13_bridge import calculate, excel_report, check_excel


def as_bytes(wb):
    out = io.BytesIO(); wb.save(out); wb.close(); return out.getvalue()


def sales_book(vendor='경기', name='검증고객', phone='01011112222', amount=100000, dt=date(2026,9,1), extra=30000):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '09월'
    for col, label in {'A':'날짜','B':'성명','D':'기종','F':'청약점','G':'개통번호','N':'판매자','P':'R/B','Q':'구두','R':'차감금액','S':'추가지원금','T':'현금개통','U':'유심비','V':'수수료','W':'정산','AE':'구두내용'}.items():
        ws[f'{col}3'] = label
    for col, value in {'A':dt,'B':name,'D':'SYNTH-S926','F':vendor,'G':phone,'N':'검증직원','P':amount,'Q':extra,'R':0,'S':700000,'T':200000,'U':7700,'V':800,'W':-999999,'AE':'추가 +3'}.items():
        ws[f'{col}4'] = value
    ws['B4'].data_type='s'
    return as_bytes(wb)


def gyeonggi_book(name='검*객', phone='010-****-2222', extra=-20000):
    wb=openpyxl.Workbook();ws=wb.active
    headers=['가입번호','판매일자','판매처','고객명','CTN','유형','유형2','모델명','액면','별도추가1','지급계','실지급계','유통망지원금','현금판매분']
    for i,h in enumerate(headers,1):ws.cell(7,i,h)
    ws.append(['SYNTH-1',date(2026,9,1),'다른 매장 계정',name,phone,'신규','','S926',-100000,extra,-100000+extra,-990000,700000,200000])
    ws['D8'].data_type='s'
    return as_bytes(wb)


def juan_book():
    wb=openpyxl.Workbook();stat=wb.active;stat.title='통계';stat['B5']='검증 계정';stat['D5']=130000
    ws=wb.create_sheet('무선')
    for c,h in {6:'①판매수수료',9:'②현금판매기기대금',10:'③미발행',11:'실지급액',12:'무선수수료상세'}.items():ws.cell(2,c,h)
    headers=['개통시간','고객명','개통번호','개통구분','모델명','무선 수수료','추가지원금','유심후청구','합계','번동료','실지급액','리베이트','별도추가']
    for i,h in enumerate(headers,1):ws.cell(3,i,h)
    ws.append([date(2026,9,1),'검*객','010-11**-*222','신규','S926',130000,700000,7700,200000,800,999999,100000,30000])
    return as_bytes(wb)


def naeway_book():
    from settlement_v13.settlement.vendors.naeway import CORE,EXCLUDE
    wb=openpyxl.Workbook();ws=wb.active
    headers=['순번','모델명','일련번호','개통판매처','고객명','개통번호','개통일','유형','총정산금','수정건',*CORE,*EXCLUDE]
    for i,h in enumerate(headers,1):ws.cell(4,i,h)
    values={'순번':1,'모델명':'S926','일련번호':'SYNTH1','개통판매처':'다른계정','고객명':'검증고객','개통번호':'01011112222','개통일':date(2026,9,1),'유형':'LG','총정산금':-999999,'리베이트':-100000,'필수팩':-30000,'추가지원금':700000,'유심':7700}
    ws.append([values.get(h) for h in headers]);return as_bytes(wb)


class V13IntegrationTest(unittest.TestCase):
    def setUp(self):
        fixtures.SignatureAndTenantTest.setUp(self)
        self.login_as_a()
        r=self.client.get('/settlements/v13');self.assertEqual(200,r.status_code)
        with self.client.session_transaction() as sess:self.token=sess['settlement_v13_csrf']

    login_as_a=fixtures.SignatureAndTenantTest.login_as_a

    def upload(self, sales=None, vendor=None, code='gyeonggi', **overrides):
        data={'csrf_token':self.token,'month':'2026-09','branch_0':str(self.a_branch),
              'sales_0':(io.BytesIO(sales or sales_book()),'sales.xlsx'),
              f'vendor_{code}':(io.BytesIO(vendor or gyeonggi_book()),'vendor.xlsx')}
        data.update(overrides)
        return self.client.post('/settlements/v13',data=data)

    def run_id(self):
        with app.app_context():return SettlementV13Run.query.one().id

    def test_original_engine_files_unchanged(self):
        root=Path(__file__).resolve().parents[1]/'settlement_v13'
        manifest=json.loads((root/'ORIGINAL_SHA256.json').read_text())
        for name, digest in manifest.items():
            if name.startswith('ui_prototype/'):
                continue
            self.assertEqual(digest,hashlib.sha256((root/name).read_bytes()).hexdigest(),name)

    def test_upload_exact_amounts_details_persist_exports_and_original(self):
        raw=sales_book();r=self.upload(sales=raw);self.assertEqual(302,r.status_code)
        rid=self.run_id()
        with app.app_context():
            run=db.session.get(SettlementV13Run,rid);result=json.loads(run.result_json)[0]
            self.assertEqual((1,1,-10000),(result['summary']['matched'],result['summary']['less'],result['summary']['total_diff']))
            item=result['items'][0];self.assertEqual((130000,120000),(item['sales_cmp'],item['vendor_cmp']))
            self.assertTrue(item['checks']);self.assertTrue(item['causes']);self.assertTrue(item['memo_match'])
            source=SettlementV13Source.query.filter_by(kind='sales').one();fid=source.id
            self.assertEqual(raw,source.content)
        page=self.client.get(r.location);self.assertEqual(200,page.status_code)
        self.assertIn(b'settlement-data',page.data);self.assertIn(b'no-store',page.headers['Cache-Control'].encode())
        self.assertEqual(raw,self.client.get(f'/settlements/v13/{rid}/files/{fid}').data)
        api=self.client.get(f'/api/settlements/v13/{rid}');self.assertEqual(result,api.json[0])
        report=self.client.get(f'/settlements/v13/{rid}/export/gyeonggi.xlsx');self.assertEqual(200,report.status_code)
        wb=openpyxl.load_workbook(io.BytesIO(report.data));self.assertTrue(any(n.startswith('001_') for n in wb.sheetnames));wb.close()
        self.assertEqual(result,self.client.get(f'/settlements/v13/{rid}/export/gyeonggi.json').json)
        with app.app_context():self.assertGreaterEqual(AuditLog.query.filter(AuditLog.action.like('정산 v1.3%')).count(),5)

    def test_all_three_vendor_adapters_keep_original_sign_and_exclusions(self):
        for code,keyword,vendor in [('gyeonggi','경기',gyeonggi_book(extra=-30000)),('juan','주안',juan_book()),('naeway','내외',naeway_book())]:
            r=self.upload(sales=sales_book(vendor=keyword),vendor=vendor,code=code);self.assertEqual(302,r.status_code,r.data)
            with app.app_context():
                result=json.loads(SettlementV13Run.query.order_by(SettlementV13Run.id.desc()).first().result_json)[0]
                self.assertEqual((1,0,130000),(result['summary']['match'],result['summary']['total_diff'],result['items'][0]['vendor_cmp']),code)

    def test_duplicate_upload_is_idempotent_duplicate_inputs_rejected(self):
        sales,vendor=sales_book(),gyeonggi_book()
        for _ in range(2):self.assertEqual(302,self.upload(sales=sales,vendor=vendor).status_code)
        with app.app_context():self.assertEqual(1,SettlementV13Run.query.count())
        r=self.upload(sales=sales,vendor=vendor,vendor_juan=(io.BytesIO(vendor),'copy.xlsx'));self.assertEqual(400,r.status_code)

    def test_other_company_and_staff_denied_on_every_data_path(self):
        self.upload();rid=self.run_id()
        with app.app_context():fid=SettlementV13Source.query.first().id
        paths=[f'/settlements/v13/{rid}',f'/api/settlements/v13/{rid}',f'/settlements/v13/{rid}/files/{fid}',f'/settlements/v13/{rid}/export/gyeonggi.xlsx',f'/settlements/v13/{rid}/export/gyeonggi.json']
        with app.app_context():u=db.session.get(User,self.a_user);u.company_code='company-b';db.session.commit()
        with self.client.session_transaction() as sess:sess['company_code']='company-b'
        for p in paths:self.assertEqual(404,self.client.get(p).status_code,p)
        self.assertEqual(403,self.upload().status_code)
        with app.app_context():u=db.session.get(User,self.a_user);u.company_code='company-a';u.role='staff';db.session.commit()
        with self.client.session_transaction() as sess:sess['company_code']='company-a'
        for p in paths+['/settlements/v13']:self.assertEqual(403,self.client.get(p).status_code,p)
        self.assertEqual(403,self.upload().status_code)

    def test_unauthenticated_pending_and_csrf(self):
        self.assertEqual(400,self.upload(csrf_token='wrong').status_code)
        with app.app_context():u=db.session.get(User,self.a_user);u.approval_pending=True;db.session.commit()
        self.assertEqual(403,self.client.get('/settlements/v13').status_code)
        with self.client.session_transaction() as sess:sess.clear()
        self.assertEqual(302,self.client.get('/settlements/v13').status_code)

    def test_invalid_vendor_year_formula_and_text_amount_do_not_save(self):
        for sales,vendor in [(sales_book(),sales_book()),(sales_book(dt=date(2025,9,1)),gyeonggi_book()),(sales_book(amount='=1+1'),gyeonggi_book()),(sales_book(amount='100,000'),gyeonggi_book())]:
            r=self.upload(sales=sales,vendor=vendor);self.assertEqual(400,r.status_code)
        with app.app_context():self.assertEqual(0,SettlementV13Run.query.count());self.assertEqual(0,SettlementV13Source.query.count())

    def test_output_escapes_html_and_excel_formula_without_changing_text(self):
        text='=1+1'
        self.upload(sales=sales_book(name=text),vendor=gyeonggi_book(name=text));rid=self.run_id()
        report=self.client.get(f'/settlements/v13/{rid}/export/gyeonggi.xlsx')
        wb=openpyxl.load_workbook(io.BytesIO(report.data))
        cells=[c for w in wb for row in w for c in row]
        self.assertFalse(any(c.data_type=='f' for c in cells));self.assertTrue(any(c.value==text for c in cells));wb.close()
        xss='</script><script>window.pwned=1</script>'
        r=self.upload(sales=sales_book(name=xss),vendor=gyeonggi_book(name=xss));self.assertEqual(302,r.status_code)
        html=self.client.get(r.location).data.decode();self.assertNotIn(xss,html);self.assertIn('\\u003c/script\\u003e',html)

    def test_backup_restore_preserves_originals_and_results(self):
        self.upload()
        with app.app_context():
            run=SettlementV13Run.query.one();original=run.result_json
            con=db.engine.raw_connection();dest=tempfile.NamedTemporaryFile(suffix='.sqlite')
            import sqlite3
            target=sqlite3.connect(dest.name)
            try:
                con.driver_connection.backup(target)
                self.assertEqual(original,target.execute('SELECT result_json FROM settlement_v13_run').fetchone()[0])
                content,digest=target.execute('SELECT content,sha256 FROM settlement_v13_source LIMIT 1').fetchone()
                self.assertEqual(digest,hashlib.sha256(content).hexdigest())
            finally:target.close();con.close();dest.close()

    def test_wrong_extension_and_incomplete_inputs(self):
        for filename,content in [('bad.exe',b'x'),('bad.xlsx',b'no zip'),('bad.xls',b'no excel')]:
            with self.assertRaises(ValueError):check_excel(content,filename)
        self.assertEqual(400,self.upload(month='2026-13').status_code)
        self.assertEqual(400,self.upload(branch_0='').status_code)


if __name__=='__main__':unittest.main()
