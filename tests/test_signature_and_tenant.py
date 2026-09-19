import os
import io
import json
import hashlib
import re
import sys
import unittest
import zipfile
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import patch

os.environ['DATABASE_URL']='sqlite:///:memory:'
os.environ['ADMIN_USERNAME']=''
os.environ['ADMIN_PASSWORD']=''

from werkzeug.security import generate_password_hash, check_password_hash
from app import app, db, User, Branch, BranchMonthlyTarget, Customer, CustomerTask, ContactLog, Booking, Inventory, InventoryMovement, DeviceMaster, Sale, SaleAddon, SaleDocument, WiredSale, CashLedger, AccountRequest, AuditLog, Price, LOGIN_STORIES, PhoneVerification, SmsCampaign, SmsCampaignRecipient, _send_sms, issue_phone_code, run_sms_campaigns, add_months, addon_rule_for_carrier, resolved_addon_rule, task_due_stage


class SignatureAndTenantTest(unittest.TestCase):
 def setUp(self):
  app.config.update(TESTING=True,SECRET_KEY='test-secret')
  self.client=app.test_client()
  with app.app_context():
   db.drop_all();db.create_all()
   a=Branch(name='A 본점',code='A99',company_code='company-a')
   a2=Branch(name='A 2호점',code='A98',company_code='company-a')
   b=Branch(name='B 본점',code='B99',company_code='company-b')
   db.session.add_all([a,a2,b]);db.session.flush()
   db.session.add_all([
    User(username='admin-a',password_hash=generate_password_hash('old-password'),role='admin',display_name='A관리자',company_code='company-a',branch_id=a.id,recovery_phone='01011112222',active=True),
    User(username='admin-b',password_hash='unused',role='admin',display_name='B관리자',company_code='company-b',branch_id=b.id,active=True),
    Customer(name='A 고객',phone='01011112222',company_code='company-a',branch_id=a.id),
    Customer(name='B 고객',phone='01033334444',company_code='company-b',branch_id=b.id)
   ]);db.session.commit()
   self.a_user=User.query.filter_by(username='admin-a').first().id
   self.a_branch=a.id;self.a2_branch=a2.id;self.b_branch=b.id
   self.b_customer=Customer.query.filter_by(name='B 고객').first().id

 def login_as_a(self):
  with self.client.session_transaction() as sess:
   sess.update(user_id=self.a_user,username='admin-a',display_name='A관리자',role='admin',company_code='company-a')

 def test_twenty_signature_scenes_exist(self):
  self.assertEqual(20,len(LOGIN_STORIES))
  response=self.client.get('/login')
  self.assertEqual(200,response.status_code)
  self.assertIn(b'SIGNATURE COLLECTION',response.data)

 def test_date_calculator_treats_leading_zero_as_day_count(self):
  self.login_as_a();response=self.client.get('/date-calculator?base=2026-09-14&days=0185')
  self.assertEqual(200,response.status_code);self.assertIn('185일'.encode(),response.data);self.assertIn('2027년 03월 18일'.encode(),response.data)

 def test_carrier_addon_rules_resolve_to_operating_defaults(self):
  self.assertEqual('M+1',addon_rule_for_carrier('SK'))
  self.assertEqual('D+95',addon_rule_for_carrier('LG U+'))
  self.assertEqual('M+3',addon_rule_for_carrier('KT'))
  self.assertEqual('D+30',resolved_addon_rule('D+30','LG'))

 def test_task_due_stage_uses_operating_color_ladder(self):
  today=date(2026,9,18)
  self.assertEqual(('due-today','D-Day'),task_due_stage(today,today))
  self.assertEqual(('due-1','D+1'),task_due_stage(today-timedelta(days=1),today))
  self.assertEqual(('due-2','D+2'),task_due_stage(today-timedelta(days=2),today))
  self.assertEqual(('due-3','D+3'),task_due_stage(today-timedelta(days=3),today))
  self.assertEqual(('due-4','D+4'),task_due_stage(today-timedelta(days=4),today))

 def test_sale_auto_addon_rule_is_saved_as_resolved_rule(self):
  self.login_as_a();opening=date(2026,1,31)
  response=self.client.post('/sales/new',data={'customer_name':'자동해지 고객','customer_phone':'01077778888','branch_id':str(self.a_branch),'opening_date':opening.isoformat(),'carrier':'SK','opening_type':'기기변경','addon_name[]':'콜키퍼','addon_rule[]':'AUTO'},follow_redirects=False)
  self.assertEqual(302,response.status_code)
  with app.app_context():
   addon=SaleAddon.query.one();self.assertEqual('M+1',addon.retention_rule);self.assertEqual(date(2026,2,28),addon.cancellation_due_date)
   task=CustomerTask.query.filter_by(task_type='부가서비스 해지').one();self.assertEqual(date(2026,2,28),task.due_date);self.assertIn('M+1',task.description)

 def test_company_customer_list_is_isolated(self):
  self.login_as_a();response=self.client.get('/customers')
  self.assertEqual(200,response.status_code)
  self.assertIn('A 고객'.encode(),response.data)
  self.assertNotIn('B 고객'.encode(),response.data)

 def test_customer_type_filter_and_export_are_company_isolated(self):
  from openpyxl import load_workbook
  with app.app_context():
   own=Customer.query.filter_by(company_code='company-a').first();own.customer_type='성지손님';own.hobbies='낚시';own.interests='인터넷 결합'
   foreign=Customer.query.execution_options(skip_tenant=True).filter_by(company_code='company-b').first();foreign.customer_type='성지손님';db.session.commit()
  self.login_as_a();page=self.client.get('/customers?customer_type=성지손님')
  self.assertEqual(200,page.status_code);self.assertIn('A 고객'.encode(),page.data);self.assertNotIn('B 고객'.encode(),page.data)
  response=self.client.get('/customers/export.xlsx?customer_type=성지손님');self.assertEqual(200,response.status_code)
  book=load_workbook(io.BytesIO(response.data),read_only=True);content=' '.join(str(cell or '') for row in book.active.iter_rows(values_only=True) for cell in row)
  self.assertIn('A 고객',content);self.assertIn('낚시',content);self.assertIn('인터넷 결합',content);self.assertNotIn('B 고객',content)

 def test_admin_notification_reports_only_own_company_downloads(self):
  with app.app_context():
   db.session.add_all([AuditLog(company_code='company-a',username='A직원',action='고객목록 다운로드',detail='성지손님 10명',ip_address='10.0.0.1'),AuditLog(company_code='company-b',username='B직원',action='고객목록 다운로드',detail='비공개 20명',ip_address='10.0.0.2')]);db.session.commit()
  self.login_as_a();response=self.client.get('/notifications');body=response.get_data(as_text=True)
  self.assertEqual(200,response.status_code);self.assertIn('A직원',body);self.assertIn('10.0.0.1',body);self.assertNotIn('B직원',body);self.assertNotIn('10.0.0.2',body)

 def test_price_excel_import_upserts_and_is_company_isolated(self):
  from openpyxl import Workbook
  with app.app_context():
   db.session.add_all([Price(device='갤럭시 A',carrier='SKT',sale_type='번호이동',price='구단가',company_code='company-a'),Price(device='타회사폰',carrier='KT',sale_type='기기변경',price='비공개',company_code='company-b')]);db.session.commit()
  book=Workbook();sheet=book.active;sheet.append(['모델명','통신사','가입유형','단가']);sheet.append(['갤럭시 A','SKT','번호이동','100,000원']);sheet.append(['아이폰 B','KT','기기변경','200,000원']);data=io.BytesIO();book.save(data);data.seek(0)
  self.login_as_a();response=self.client.post('/prices/import',data={'file':(data,'partner.xlsx')},content_type='multipart/form-data',follow_redirects=True);body=response.get_data(as_text=True)
  self.assertEqual(200,response.status_code);self.assertIn('신규 1건, 수정 1건'.encode(),response.data);self.assertIn('아이폰 B',body);self.assertNotIn('타회사폰',body)
  with app.app_context():
   self.assertEqual('100,000원',Price.query.execution_options(skip_tenant=True).filter_by(company_code='company-a',device='갤럭시 A').one().price)
   self.assertEqual(2,Price.query.execution_options(skip_tenant=True).filter_by(company_code='company-a').count())

 def test_price_management_rejects_staff(self):
  self.login_as_a()
  with self.client.session_transaction() as sess:sess['role']='staff'
  with app.app_context():
   user=db.session.get(User,self.a_user);user.role='staff';db.session.commit()
  self.assertEqual(403,self.client.get('/prices').status_code)

 def test_sale_form_receives_only_company_price_policies(self):
  with app.app_context():
   db.session.add_all([Price(device='갤럭시 S26',carrier='SK',sale_type='번호이동',price='500,000원',rebate_amount=500000,company_code='company-a'),Price(device='타회사 비공개폰',carrier='KT',sale_type='기기변경',price='900,000원',rebate_amount=900000,company_code='company-b')]);db.session.commit()
  self.login_as_a();response=self.client.get('/sales/new');body=response.get_data(as_text=True)
  self.assertEqual(200,response.status_code);self.assertIn('"rebate_amount": 500000',body);self.assertIn('단가표 자동 적용',body);self.assertNotIn('"rebate_amount": 900000',body)

 def test_sale_form_provides_active_device_autocomplete(self):
  with app.app_context():
   db.session.add_all([DeviceMaster(manufacturer='삼성',model='갤럭시 S26 울트라',capacities='256GB,512GB',colors='블랙,화이트',active=True),DeviceMaster(manufacturer='애플',model='숨김 아이폰',active=False)]);db.session.commit()
  self.login_as_a();response=self.client.get('/sales/new');body=response.get_data(as_text=True)
  self.assertEqual(200,response.status_code);match=re.search(r'const saleDeviceMaster=(.*?);',body);self.assertIsNotNone(match);devices=json.loads(match.group(1));models=[x['model'] for x in devices]
  self.assertIn('갤럭시 S26 울트라',models);self.assertNotIn('숨김 아이폰',models);self.assertEqual('256GB,512GB',next(x['capacities'] for x in devices if x['model']=='갤럭시 S26 울트라'));self.assertIn('applyDeviceMaster',body)

 def test_performance_report_includes_staff_workload_and_blocks_foreign_tasks(self):
  with app.app_context():
   own_sale=Sale(customer_name='A 업무고객',opening_date=date.today(),branch_id=self.a_branch,assigned_staff='A관리자');foreign_sale=Sale(customer_name='B 업무고객',opening_date=date.today(),branch_id=self.b_branch,assigned_staff='B관리자');db.session.add_all([own_sale,foreign_sale]);db.session.flush()
   db.session.add_all([CustomerTask(sale_id=own_sale.id,task_type='고객약속',title='A 공개업무',due_date=date.today()-timedelta(days=1),assigned_staff='A관리자',status='처리예정'),CustomerTask(sale_id=foreign_sale.id,task_type='고객약속',title='B 비공개업무',due_date=date.today()-timedelta(days=1),assigned_staff='B관리자',status='처리예정')]);db.session.commit()
  self.login_as_a();response=self.client.get(f'/reports/sales-performance?month={date.today():%Y-%m}');body=response.get_data(as_text=True)
  self.assertEqual(200,response.status_code);self.assertIn('미처리',body);self.assertIn('기한초과',body);self.assertIn('A관리자',body);self.assertNotIn('B관리자',body)

 def test_contact_failure_creates_customer_history_once(self):
  with app.app_context():
   customer=Customer.query.filter_by(company_code='company-a').first();customer_id=customer.id;task=CustomerTask(customer_id=customer_id,task_type='OB 재연락',title='고객 재통화',due_date=date.today(),status='처리예정');db.session.add(task);db.session.commit();task_id=task.id
  self.login_as_a();response=self.client.post(f'/tasks/{task_id}/status',data={'status':'연락안됨'},follow_redirects=False);self.assertEqual(302,response.status_code)
  self.client.post(f'/tasks/{task_id}/status',data={'status':'연락안됨'},follow_redirects=False)
  with app.app_context():
   logs=ContactLog.query.filter_by(customer_id=customer_id,outcome='연락안됨').all();self.assertEqual(1,len(logs));self.assertIn('카카오채널',logs[0].note)

 def test_booking_is_store_scoped_and_shown_on_dashboard(self):
  visit=f'{date.today().isoformat()}T14:30'
  self.login_as_a();response=self.client.post('/bookings',data={'branch_id':str(self.a_branch),'name':'A 예약고객','phone':'010-7777-8888','visit_date':visit,'device':'갤럭시 S26'},follow_redirects=False)
  self.assertEqual(302,response.status_code)
  with app.app_context():
   own=Booking.query.execution_options(skip_tenant=True).filter_by(name='A 예약고객').one();self.assertEqual(('company-a',self.a_branch,'01077778888','예약'),(own.company_code,own.branch_id,own.phone,own.status))
   db.session.add_all([Booking(name='A 타지점예약',visit_date=visit,company_code='company-a',branch_id=self.a2_branch),Booking(name='B 비공개예약',visit_date=visit,company_code='company-b',branch_id=self.b_branch)]);db.session.commit()
  dashboard=self.client.get('/');self.assertIn('A 예약고객'.encode(),dashboard.data);self.assertIn('A 타지점예약'.encode(),dashboard.data);self.assertNotIn('B 비공개예약'.encode(),dashboard.data)
  with app.app_context():user=db.session.get(User,self.a_user);user.role='staff';user.branch_id=self.a_branch;db.session.commit()
  listing=self.client.get('/bookings');self.assertIn('A 예약고객'.encode(),listing.data);self.assertNotIn('A 타지점예약'.encode(),listing.data);self.assertNotIn('B 비공개예약'.encode(),listing.data)
  with app.app_context():foreign_store=Booking.query.execution_options(skip_tenant=True).filter_by(name='A 타지점예약').one().id
  self.assertEqual(403,self.client.post(f'/bookings/{foreign_store}/status',data={'status':'방문완료'}).status_code)

 def test_deactivated_user_session_is_revoked_on_next_request(self):
  self.login_as_a()
  with app.app_context():
   user=db.session.get(User,self.a_user);user.active=False;db.session.commit()
  response=self.client.get('/',follow_redirects=False)
  self.assertEqual(302,response.status_code);self.assertTrue(response.headers['Location'].endswith('/login'))
  with self.client.session_transaction() as sess:self.assertNotIn('user_id',sess)

 def test_role_and_branch_changes_refresh_in_existing_session(self):
  self.login_as_a()
  with app.app_context():
   user=db.session.get(User,self.a_user);user.role='staff';user.branch_id=self.a2_branch;user.display_name='변경직원';db.session.commit()
  response=self.client.get('/',follow_redirects=False)
  self.assertEqual(200,response.status_code)
  with self.client.session_transaction() as sess:
   self.assertEqual(('staff',self.a2_branch,'변경직원'),(sess['role'],sess['branch_id'],sess['display_name']))

 def test_staff_choices_and_direct_account_access_are_company_isolated(self):
  self.login_as_a()
  form=self.client.get('/sales/new')
  self.assertIn('A관리자'.encode(),form.data);self.assertNotIn('B관리자'.encode(),form.data)
  with app.app_context():foreign=User.query.execution_options(skip_tenant=True).filter_by(username='admin-b').one().id
  self.assertEqual(404,self.client.get(f'/staff/{foreign}/edit').status_code)

 def test_admin_sales_and_inventory_lists_are_company_isolated(self):
  with app.app_context():
   db.session.add_all([
    Sale(customer_name='A 판매고객',customer_phone='01011112222',opening_date=date.today(),device='A전용단말',serial_number='A-SERIAL',branch_id=self.a_branch),
    Sale(customer_name='B 비공개판매',customer_phone='01033334444',opening_date=date.today(),device='B전용단말',serial_number='B-SERIAL',branch_id=self.b_branch),
    Inventory(serial_number='A-STOCK',model='A재고모델',branch_id=self.a_branch,status='보유중'),
    Inventory(serial_number='B-STOCK',model='B비공개재고',branch_id=self.b_branch,status='보유중')
   ]);db.session.commit()
  self.login_as_a()
  sales=self.client.get('/sales');stock=self.client.get('/inventory')
  self.assertIn('A 판매고객'.encode(),sales.data);self.assertNotIn('B 비공개판매'.encode(),sales.data)
  self.assertIn(b'A-STOCK',stock.data);self.assertNotIn(b'B-STOCK',stock.data)

 def test_inventory_excel_import_infers_device_and_skips_duplicates(self):
  from openpyxl import Workbook
  with app.app_context():
   db.session.add(DeviceMaster(manufacturer='삼성',model='갤럭시 S26',capacities='256GB,512GB',colors='블랙,화이트',active=True));db.session.add(Inventory(serial_number='DUP-001',model='기존재고',branch_id=self.a_branch,status='보유중'));db.session.commit()
  book=Workbook();sheet=book.active;sheet.append(['시리얼번호','기종','통신사','매입가','보관위치']);sheet.append(['NEW-001','갤럭시 S26 256GB 블랙','SKT','1,250,000원','창고 A']);sheet.append(['DUP-001','갤럭시 S26','KT',900000,'창고 B']);data=io.BytesIO();book.save(data);data.seek(0)
  self.login_as_a();response=self.client.post('/inventory/import',data={'branch_id':str(self.a_branch),'file':(data,'stock.xlsx')},content_type='multipart/form-data',follow_redirects=True)
  self.assertEqual(200,response.status_code);self.assertIn('등록 1건, 중복 1건'.encode(),response.data)
  with app.app_context():
   item=Inventory.query.filter_by(serial_number='NEW-001').one();self.assertEqual(('갤럭시 S26','삼성','256GB','블랙',1250000,self.a_branch),(item.model,item.manufacturer,item.capacity,item.color,item.purchase_price,item.branch_id));self.assertEqual('엑셀입고',InventoryMovement.query.filter_by(inventory_id=item.id).one().action)

 def test_inventory_excel_import_rejects_staff(self):
  self.login_as_a()
  with app.app_context():user=db.session.get(User,self.a_user);user.role='staff';db.session.commit()
  self.assertEqual(403,self.client.post('/inventory/import',data={}).status_code)

 def test_admin_cannot_delete_other_company_sale(self):
  with app.app_context():
   foreign=Sale(customer_name='B 삭제차단',opening_date=date.today(),branch_id=self.b_branch);db.session.add(foreign);db.session.commit();sale_id=foreign.id
  self.login_as_a();response=self.client.post(f'/sales/{sale_id}/delete')
  self.assertEqual(403,response.status_code)
  with app.app_context():self.assertIsNotNone(Sale.query.execution_options(skip_tenant=True).filter_by(id=sale_id).first())

 def test_admin_wired_and_cash_lists_are_company_isolated(self):
  with app.app_context():
   db.session.add_all([WiredSale(sale_date=date.today(),customer_name='A 유선고객',branch_id=self.a_branch),WiredSale(sale_date=date.today(),customer_name='B 비공개유선',branch_id=self.b_branch),CashLedger(ledger_date=date.today(),branch_id=self.a_branch,direction='입금',category='A시재',amount=1000),CashLedger(ledger_date=date.today(),branch_id=self.b_branch,direction='입금',category='B비공개시재',amount=2000)]);db.session.commit()
  self.login_as_a();wired=self.client.get('/wired-sales');ledger=self.client.get('/cash-ledger')
  self.assertIn('A 유선고객'.encode(),wired.data);self.assertNotIn('B 비공개유선'.encode(),wired.data)
  self.assertIn('A시재'.encode(),ledger.data);self.assertNotIn('B비공개시재'.encode(),ledger.data)

 def test_monthly_performance_graph_groups_stores_and_staff(self):
  with app.app_context():
   db.session.add_all([
    Sale(customer_name='A 실적고객',opening_date=date.today(),branch_id=self.a_branch,assigned_staff='김판매',settlement_amount_v2=500000,final_margin=120000),
    WiredSale(sale_date=date.today(),customer_name='A 유선실적',branch_id=self.a_branch,assigned_staff='김판매',settlement_amount=300000,final_margin=80000),
    Sale(customer_name='B 비공개실적',opening_date=date.today(),branch_id=self.b_branch,assigned_staff='타회사직원',settlement_amount_v2=900000,final_margin=400000)
   ]);db.session.commit()
  self.login_as_a();response=self.client.get(f'/reports/sales-performance?month={date.today():%Y-%m}')
  self.assertEqual(200,response.status_code);body=response.get_data(as_text=True)
  self.assertIn('A 본점',body);self.assertIn('김판매',body);self.assertIn('800,000원',body);self.assertIn('200,000원',body)
  self.assertNotIn('B 비공개실적',body);self.assertNotIn('타회사직원',body)

 def test_admin_sets_company_scoped_branch_targets_and_sees_attainment(self):
  month=date.today().strftime('%Y-%m');self.login_as_a();response=self.client.post('/reports/sales-performance/targets',data={'month':month,'branch_id':[str(self.a_branch)],f'mobile_goal_{self.a_branch}':'10',f'wired_goal_{self.a_branch}':'5',f'margin_goal_{self.a_branch}':'1,000,000'},follow_redirects=True)
  self.assertEqual(200,response.status_code);body=response.get_data(as_text=True);self.assertIn('15건',body);self.assertIn('실적목표를 저장',body)
  with app.app_context():
   target=BranchMonthlyTarget.query.execution_options(skip_tenant=True).filter_by(company_code='company-a',branch_id=self.a_branch,month=month).one();self.assertEqual((10,5,1000000),(target.mobile_goal,target.wired_goal,target.margin_goal))
   self.assertIsNone(BranchMonthlyTarget.query.execution_options(skip_tenant=True).filter_by(company_code='company-b',branch_id=self.a_branch,month=month).first())

 def test_branch_target_rejects_another_company_branch(self):
  month=date.today().strftime('%Y-%m');self.login_as_a();response=self.client.post('/reports/sales-performance/targets',data={'month':month,'branch_id':[str(self.b_branch)],f'mobile_goal_{self.b_branch}':'99'})
  self.assertEqual(302,response.status_code)
  with app.app_context():self.assertIsNone(BranchMonthlyTarget.query.execution_options(skip_tenant=True).filter_by(branch_id=self.b_branch,month=month).first())

 def test_sales_settlement_filters_and_admin_status_update(self):
  with app.app_context():
   own=Sale(customer_name='정산대상고객',opening_date=date.today(),branch_id=self.a_branch,assigned_staff='김정산',settlement_amount_v2=500000,final_margin=120000,settlement_status='미지급');foreign=Sale(customer_name='타회사정산',opening_date=date.today(),branch_id=self.b_branch,assigned_staff='타직원',settlement_status='추가금');db.session.add_all([own,foreign]);db.session.commit();own_id=own.id
  self.login_as_a();page=self.client.get(f'/sales?month={date.today():%Y-%m}&settlement_status=미지급&staff=김정산');body=page.get_data(as_text=True)
  self.assertEqual(200,page.status_code);self.assertIn('정산대상고객',body);self.assertIn('미확인 정산',body);self.assertNotIn('타회사정산',body)
  response=self.client.post(f'/sales/{own_id}/settlement-status',data={'status':'정상'},follow_redirects=True);self.assertEqual(200,response.status_code)
  with app.app_context():
   item=db.session.get(Sale,own_id);self.assertEqual('정상',item.settlement_status);self.assertEqual('A관리자',item.settlement_checked_by);self.assertIsNotNone(item.settlement_checked_at)

 def test_overdue_settlement_alerts_are_company_scoped(self):
  old=date.today()-timedelta(days=5)
  with app.app_context():db.session.add_all([Sale(customer_name='A 지연정산',opening_date=old,branch_id=self.a_branch,assigned_staff='A관리자',settlement_status='미지급'),Sale(customer_name='A 정상완료',opening_date=old,branch_id=self.a_branch,settlement_status='정상'),Sale(customer_name='B 비공개지연',opening_date=old,branch_id=self.b_branch,settlement_status='추가금')]);db.session.commit()
  self.login_as_a()
  for path in ['/', '/notifications', '/sales?settlement_age=overdue']:
   response=self.client.get(path);body=response.get_data(as_text=True);self.assertEqual(200,response.status_code);self.assertIn('A 지연정산',body);self.assertNotIn('A 정상완료',body);self.assertNotIn('B 비공개지연',body)
  with self.client.session_transaction() as sess:self.assertGreaterEqual(sess.get('user_id',0),1)

 def test_monthly_performance_excel_is_company_isolated(self):
  from openpyxl import load_workbook
  with app.app_context():
   db.session.add_all([Sale(customer_name='A 보고서',opening_date=date.today(),branch_id=self.a_branch,assigned_staff='A직원',settlement_amount_v2=100000,final_margin=30000),Sale(customer_name='B 보고서',opening_date=date.today(),branch_id=self.b_branch,assigned_staff='B직원',settlement_amount_v2=700000,final_margin=200000)]);db.session.commit()
  self.login_as_a();response=self.client.get(f'/reports/sales-performance.xlsx?month={date.today():%Y-%m}')
  self.assertEqual(200,response.status_code);book=load_workbook(io.BytesIO(response.data),read_only=True)
  content=' '.join(str(cell or '') for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for cell in row)
  self.assertIn('A직원',content);self.assertNotIn('B직원',content);self.assertEqual(['매장별 실적','직원별 실적'],book.sheetnames)

 def test_sales_settlement_excel_keeps_filters_and_company_scope(self):
  from openpyxl import load_workbook
  with app.app_context():
   db.session.add_all([Sale(customer_name='A 미지급엑셀',opening_date=date.today(),branch_id=self.a_branch,assigned_staff='A정산',settlement_amount_v2=110000,final_margin=22000,settlement_status='미지급'),Sale(customer_name='A 정상제외',opening_date=date.today(),branch_id=self.a_branch,assigned_staff='A정산',settlement_status='정상'),Sale(customer_name='B 엑셀비공개',opening_date=date.today(),branch_id=self.b_branch,assigned_staff='B정산',settlement_status='미지급')]);db.session.commit()
  self.login_as_a();response=self.client.get(f'/sales/export.xlsx?month={date.today():%Y-%m}&settlement_status=미지급&staff=A정산')
  self.assertEqual(200,response.status_code);book=load_workbook(io.BytesIO(response.data),read_only=True);content=' '.join(str(cell or '') for row in book.active.iter_rows(values_only=True) for cell in row)
  self.assertIn('A 미지급엑셀',content);self.assertNotIn('A 정상제외',content);self.assertNotIn('B 엑셀비공개',content);self.assertIn('110000',content)

 def test_ob_management_uses_latest_sale_and_customer_type_filter(self):
  old=add_months(date.today(),-31);recent=add_months(date.today(),-2)
  with app.app_context():
   target=Customer(name='장기 성지고객',phone='01055551111',company_code='company-a',branch_id=self.a_branch,customer_type='성지손님')
   renewed=Customer(name='최근 재개통고객',phone='01055552222',company_code='company-a',branch_id=self.a_branch,customer_type='성지손님')
   foreign=Customer(name='타회사 장기고객',phone='01055553333',company_code='company-b',branch_id=self.b_branch,customer_type='성지손님')
   db.session.add_all([target,renewed,foreign,Sale(customer_name=target.name,customer_phone=target.phone,opening_date=old,branch_id=self.a_branch),Sale(customer_name=renewed.name,customer_phone=renewed.phone,opening_date=old,branch_id=self.a_branch),Sale(customer_name=renewed.name,customer_phone=renewed.phone,opening_date=recent,branch_id=self.a_branch),Sale(customer_name=foreign.name,customer_phone=foreign.phone,opening_date=old,branch_id=self.b_branch)]);db.session.commit()
  self.login_as_a();response=self.client.get('/ob-management?age=30&customer_type=성지손님')
  self.assertEqual(200,response.status_code);body=response.get_data(as_text=True)
  self.assertIn('장기 성지고객',body);self.assertNotIn('최근 재개통고객',body);self.assertNotIn('타회사 장기고객',body)

 def test_ob_management_excel_contains_filtered_customers(self):
  from openpyxl import load_workbook
  with app.app_context():
   customer=Customer(name='OB 엑셀고객',phone='01066661111',company_code='company-a',branch_id=self.a_branch,customer_type='로드손님');db.session.add(customer);db.session.add(Sale(customer_name=customer.name,customer_phone=customer.phone,opening_date=add_months(date.today(),-20),branch_id=self.a_branch,assigned_staff='담당직원'));db.session.commit()
  self.login_as_a();response=self.client.get('/ob-management.xlsx?age=18&customer_type=로드손님')
  self.assertEqual(200,response.status_code);book=load_workbook(io.BytesIO(response.data),read_only=True);content=' '.join(str(cell or '') for row in book.active.iter_rows(values_only=True) for cell in row)
  self.assertIn('OB 엑셀고객',content);self.assertIn('로드손님',content);self.assertIn('담당직원',content)

 def test_customer_profile_saves_hobbies_and_interests(self):
  self.login_as_a();response=self.client.post('/customers/new',data={'name':'관심고객','phone':'01077771111','branch_id':self.a_branch,'customer_type':'성지손님','hobbies':'낚시, 골프','interests':'카메라, 인터넷 결합'},follow_redirects=False)
  self.assertEqual(302,response.status_code)
  with app.app_context():
   customer=Customer.query.filter_by(name='관심고객').one();self.assertEqual(('성지손님','낚시, 골프','카메라, 인터넷 결합'),(customer.customer_type,customer.hobbies,customer.interests))

 def test_admin_can_bulk_import_customers_without_phone_duplicates(self):
  from openpyxl import Workbook
  book=Workbook();sheet=book.active;sheet.append(['고객명','휴대전화','고객유형','통신사','취미','문자수신동의'])
  sheet.append(['신규고객',1077778888,'성지손님','SKT','낚시','동의']);sheet.append(['중복고객','01011112222','기존손님','KT','','미동의']);sheet.append(['번호오류','123','로드손님','LG','','동의'])
  payload=io.BytesIO();book.save(payload);payload.seek(0);self.login_as_a()
  response=self.client.post('/customers/import',data={'branch_id':str(self.a_branch),'file':(payload,'customers.xlsx')},content_type='multipart/form-data',follow_redirects=True)
  self.assertEqual(200,response.status_code);self.assertIn('신규 1명, 중복 1명, 제외 1명'.encode(),response.data)
  with app.app_context():
   customer=Customer.query.filter_by(company_code='company-a',phone='01077778888').one();self.assertEqual(('신규고객','성지손님','SK','낚시',True,self.a_branch),(customer.name,customer.customer_type,customer.carrier,customer.hobbies,customer.marketing_consent,customer.branch_id));self.assertIsNone(Customer.query.filter_by(name='중복고객').first())

 def test_customer_bulk_import_is_admin_only(self):
  self.login_as_a()
  with app.app_context():user=db.session.get(User,self.a_user);user.role='staff';db.session.commit()
  self.assertEqual(403,self.client.get('/customers/template.xlsx').status_code);self.assertEqual(403,self.client.post('/customers/import').status_code)

 def test_next_contact_date_creates_single_followup_task(self):
  followup=date.today()+timedelta(days=3)
  with app.app_context():
   customer=Customer.query.filter_by(company_code='company-a').first();db.session.add(Sale(customer_name=customer.name,customer_phone=customer.phone,opening_date=date.today(),branch_id=self.a_branch));db.session.commit();customer_id=customer.id
  self.login_as_a();payload={'channel':'전화','outcome':'재통화','note':'요금제 변경 재안내','next_contact_date':followup.isoformat(),'branch_id':self.a_branch}
  self.client.post(f'/customers/{customer_id}/contact-log',data=payload);self.client.post(f'/customers/{customer_id}/contact-log',data=payload)
  with app.app_context():
   tasks=CustomerTask.query.filter_by(customer_id=customer_id,task_type='OB 재연락',due_date=followup).all();self.assertEqual(1,len(tasks));self.assertEqual('요금제 변경 재안내',tasks[0].description)

 def test_admin_backup_excludes_other_company_records(self):
  from openpyxl import load_workbook
  with app.app_context():
   db.session.add_all([Sale(customer_name='A 백업고객',opening_date=date.today(),branch_id=self.a_branch),Sale(customer_name='B 백업비공개',opening_date=date.today(),branch_id=self.b_branch),Inventory(serial_number='A-BACKUP',model='A모델',branch_id=self.a_branch),Inventory(serial_number='B-BACKUP',model='B모델',branch_id=self.b_branch),Booking(name='A 예약백업',visit_date=f'{date.today()}T10:00',company_code='company-a',branch_id=self.a_branch),Booking(name='B 예약비공개',visit_date=f'{date.today()}T11:00',company_code='company-b',branch_id=self.b_branch)]);db.session.commit()
  self.login_as_a();response=self.client.get('/admin/backup.xlsx')
  self.assertEqual(200,response.status_code);book=load_workbook(io.BytesIO(response.data),read_only=True)
  sales=' '.join(str(cell or '') for row in book['판매일보'].iter_rows(values_only=True) for cell in row);stock=' '.join(str(cell or '') for row in book['재고'].iter_rows(values_only=True) for cell in row);bookings=' '.join(str(cell or '') for row in book['방문예약'].iter_rows(values_only=True) for cell in row)
  self.assertEqual(11,len(book.sheetnames));self.assertIn('A 백업고객',sales);self.assertNotIn('B 백업비공개',sales);self.assertIn('A-BACKUP',stock);self.assertNotIn('B-BACKUP',stock);self.assertIn('A 예약백업',bookings);self.assertNotIn('B 예약비공개',bookings)

 def test_complete_archive_contains_documents_and_verified_manifest(self):
  with app.app_context():
   sale=Sale(customer_name='A 전체백업고객',opening_date=date.today(),branch_id=self.a_branch);other=Sale(customer_name='B 전체백업비공개',opening_date=date.today(),branch_id=self.b_branch)
   db.session.add_all([sale,other]);db.session.flush();db.session.add_all([SaleDocument(sale_id=sale.id,branch_id=self.a_branch,original_name='계약서.pdf',content_type='application/pdf',file_size=8,file_data=b'%PDF-safe',uploaded_by='A관리자'),SaleDocument(sale_id=other.id,branch_id=self.b_branch,original_name='비공개.pdf',content_type='application/pdf',file_size=6,file_data=b'secret',uploaded_by='B관리자')]);db.session.commit()
  self.login_as_a();response=self.client.get('/admin/backup.zip')
  self.assertEqual(200,response.status_code);self.assertEqual('application/zip',response.mimetype)
  with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
   self.assertIsNone(archive.testzip());manifest=json.loads(archive.read('manifest.json'));sales=json.loads(archive.read('data/sales.json'))
   self.assertEqual('trustflow-complete-backup',manifest['format']);self.assertEqual('company-a',manifest['company_code'])
   self.assertIn('A 전체백업고객',{row['customer_name'] for row in sales});self.assertNotIn('B 전체백업비공개',{row['customer_name'] for row in sales})
   document_paths=[name for name in archive.namelist() if name.startswith('documents/')]
   self.assertEqual(1,len(document_paths));self.assertEqual(b'%PDF-safe',archive.read(document_paths[0]))
   for path,digest in manifest['files'].items():self.assertEqual(hashlib.sha256(archive.read(path)).hexdigest(),digest)

 def test_direct_cross_company_customer_access_is_blocked(self):
  self.login_as_a();response=self.client.get(f'/customers/{self.b_customer}')
  self.assertEqual(404,response.status_code)

 def test_login_fields_are_in_company_user_password_order(self):
  body=self.client.get('/login').get_data(as_text=True)
  self.assertLess(body.index('name="company_code"'),body.index('name="username"'))
  self.assertLess(body.index('name="username"'),body.index('name="password"'))

 def test_find_id_requires_phone_code(self):
  response=self.client.post('/find-id',data={'action':'send','company_code':'company-a','display_name':'A관리자','phone':'01011112222'})
  self.assertIn('인증번호를 문자로 보냈습니다'.encode(),response.data)
  response=self.client.post('/find-id',data={'action':'verify','company_code':'company-a','display_name':'A관리자','phone':'01011112222','code':'123456'})
  self.assertIn(b'admin-a',response.data)

 def test_password_can_reset_after_phone_code(self):
  base={'company_code':'company-a','username':'admin-a','display_name':'A관리자','phone':'01011112222'}
  self.client.post('/password-help',data={**base,'action':'send'})
  response=self.client.post('/password-help',data={**base,'action':'reset','code':'123456','new_password':'new-password'})
  self.assertIn('비밀번호 변경 완료'.encode(),response.data)
  with app.app_context():self.assertTrue(check_password_hash(User.query.filter_by(username='admin-a').first().password_hash,'new-password'))

 def test_password_reset_rejects_wrong_code(self):
  base={'company_code':'company-a','username':'admin-a','display_name':'A관리자','phone':'01011112222'}
  self.client.post('/password-help',data={**base,'action':'send'})
  response=self.client.post('/password-help',data={**base,'action':'reset','code':'999999','new_password':'new-password'})
  self.assertIn('인증번호가 올바르지 않습니다'.encode(),response.data)

 def test_sms_request_is_limited_per_phone_and_hour(self):
  with app.app_context():
   for purpose in ('one','two','three','four','five'):
    ok,_=issue_phone_code(purpose,'company-a','010-1111-2222');self.assertTrue(ok)
   ok,message=issue_phone_code('six','company-a','01011112222')
   self.assertFalse(ok);self.assertIn('1시간 후',message)

 def test_sms_resend_cooldown_keeps_code_input_visible(self):
  base={'company_code':'company-a','display_name':'A관리자','phone':'010-1111-2222','action':'send'}
  self.client.post('/find-id',data=base)
  response=self.client.post('/find-id',data=base)
  self.assertIn('1분 후 다시 요청'.encode(),response.data);self.assertIn('name="code"'.encode(),response.data)

 def test_signup_requires_phone_verification_and_creates_approval_request(self):
  base={'company_code':'company-a','display_name':'신입직원','phone':'010-5555-6666','username':'new-staff','password':'safe-password'}
  response=self.client.post('/signup',data={**base,'action':'register','code':'123456'})
  self.assertIn('인증번호가 만료'.encode(),response.data)
  with app.app_context():self.assertIsNone(User.query.filter_by(username='new-staff').first())
  self.client.post('/signup',data={**base,'action':'send'})
  response=self.client.post('/signup',data={**base,'action':'register','code':'123456'},follow_redirects=False)
  self.assertEqual(302,response.status_code)
  with app.app_context():
   user=User.query.filter_by(company_code='company-a',username='new-staff').one();request_item=AccountRequest.query.filter_by(company_code='company-a',username='new-staff').one()
   self.assertFalse(user.active);self.assertEqual('01055556666',user.recovery_phone);self.assertEqual(('회원가입','대기'),(request_item.request_type,request_item.status))

 def test_admin_approval_activates_signup_account(self):
  with app.app_context():
   user=User(username='pending-staff',password_hash=generate_password_hash('safe-password'),role='staff',display_name='대기직원',company_code='company-a',recovery_phone='01055556666',active=False)
   db.session.add(user);db.session.add(AccountRequest(request_type='회원가입',company_code='company-a',username='pending-staff',display_name='대기직원',phone='01055556666',status='대기'));db.session.commit();request_id=AccountRequest.query.filter_by(username='pending-staff').one().id
  self.login_as_a()
  with patch('app._send_sms',return_value=True) as sms:
   response=self.client.post(f'/account-requests/{request_id}/complete',data={'decision':'approve','branch_id':self.a_branch},follow_redirects=False)
   sms.assert_called_once();self.assertEqual('01055556666',sms.call_args.args[0]);self.assertIn('가입이 승인',sms.call_args.args[1])
  self.assertEqual(302,response.status_code)
  with app.app_context():
   approved=User.query.filter_by(company_code='company-a',username='pending-staff').one()
   self.assertTrue(approved.active);self.assertEqual(self.a_branch,approved.branch_id)
   self.assertEqual('승인',AccountRequest.query.get(request_id).status)

 def test_admin_cannot_approve_another_company_signup(self):
  with app.app_context():
   db.session.add(User(username='other-pending',password_hash='unused',company_code='company-b',active=False));db.session.add(AccountRequest(request_type='회원가입',company_code='company-b',username='other-pending',status='대기'));db.session.commit();request_id=AccountRequest.query.filter_by(username='other-pending').one().id
  self.login_as_a();response=self.client.post(f'/account-requests/{request_id}/complete',data={'decision':'approve'})
  self.assertEqual(404,response.status_code)
  with app.app_context():self.assertFalse(User.query.execution_options(skip_tenant=True).filter_by(company_code='company-b',username='other-pending').one().active)

 def test_signup_approval_requires_company_branch(self):
  with app.app_context():
   db.session.add(User(username='branchless',password_hash='unused',company_code='company-a',active=False));db.session.add(AccountRequest(request_type='회원가입',company_code='company-a',username='branchless',status='대기'));db.session.commit();request_id=AccountRequest.query.filter_by(username='branchless').one().id
  self.login_as_a()
  response=self.client.post(f'/account-requests/{request_id}/complete',data={'decision':'approve'},follow_redirects=True)
  self.assertIn('소속 지점을 선택'.encode(),response.data)
  response=self.client.post(f'/account-requests/{request_id}/complete',data={'decision':'approve','branch_id':self.b_branch})
  self.assertEqual(403,response.status_code)
  with app.app_context():self.assertFalse(User.query.filter_by(company_code='company-a',username='branchless').one().active)

 def test_admin_notification_includes_only_own_company_signup_requests(self):
  with app.app_context():
   db.session.add_all([AccountRequest(request_type='회원가입',company_code='company-a',username='a-pending',display_name='A신청자',phone='01011110000',status='대기'),AccountRequest(request_type='회원가입',company_code='company-b',username='b-pending',display_name='B신청자',phone='01022220000',status='대기')]);db.session.commit()
  self.login_as_a();response=self.client.get('/notifications')
  self.assertEqual(200,response.status_code);self.assertIn('A신청자'.encode(),response.data);self.assertNotIn('B신청자'.encode(),response.data)
  self.assertIn(b'>1</b>',response.data)

 def test_login_is_limited_after_five_failures(self):
  for _ in range(5):self.client.post('/login',data={'company_code':'company-a','username':'admin-a','password':'wrong'})
  response=self.client.post('/login',data={'company_code':'company-a','username':'admin-a','password':'old-password'})
  self.assertEqual(429,response.status_code)

 def test_same_personal_id_can_exist_in_different_companies(self):
  with app.app_context():
   db.session.add(User(username='shared-id',password_hash=generate_password_hash('password-a'),company_code='company-a',active=True))
   db.session.add(User(username='shared-id',password_hash=generate_password_hash('password-b'),company_code='company-b',active=True));db.session.commit()
   self.assertEqual(2,User.query.execution_options(skip_tenant=True).filter_by(username='shared-id').count())

 def test_solapi_sms_uses_registered_sender_and_normalized_numbers(self):
  captured={}
  class FakeService:
   def __init__(self,api_key,api_secret):captured.update(api_key=api_key,api_secret=api_secret)
   def send(self,message):captured['message']=message;return SimpleNamespace(group_info=SimpleNamespace(count=SimpleNamespace(registered_failed=0)))
  class FakeMessage:
   def __init__(self,**kwargs):self.__dict__.update(kwargs)
  fake_solapi=SimpleNamespace(SolapiMessageService=FakeService)
  fake_model=SimpleNamespace(RequestMessage=FakeMessage)
  with patch.dict(os.environ,{'SOLAPI_API_KEY':'key','SOLAPI_API_SECRET':'secret','SMS_SENDER':'010-7667-1100'},clear=False), patch.dict(sys.modules,{'solapi':fake_solapi,'solapi.model':fake_model}):
   app.config['TESTING']=False
   try:self.assertTrue(_send_sms('010-1234-5678','인증번호 테스트'))
   finally:app.config['TESTING']=True
  self.assertEqual('01076671100',captured['message'].from_)
  self.assertEqual('01012345678',captured['message'].to)

 def test_campaign_sends_only_consented_customers_within_daily_limit(self):
  with app.app_context():
   consented=Customer.query.filter_by(company_code='company-a').first();consented.marketing_consent=True
   other=Customer(name='A 수신동의2',phone='01077778888',company_code='company-a',branch_id=self.a_branch,marketing_consent=True)
   campaign=SmsCampaign(company_code='company-a',name='행사',message='혜택 안내',start_month='2026-09',end_month='2026-09',daily_limit=1,status='진행중',sender_name='고무신모바일',sender_contact='01076671100',opt_out_number='0800000000')
   db.session.add_all([other,campaign]);db.session.flush()
   db.session.add_all([SmsCampaignRecipient(campaign_id=campaign.id,customer_id=consented.id,phone=consented.phone,scheduled_date=date.today()),SmsCampaignRecipient(campaign_id=campaign.id,customer_id=other.id,phone=other.phone,scheduled_date=date.today())]);db.session.commit()
  with app.app_context(),patch('app._send_sms',return_value=True) as sms:
   result=run_sms_campaigns(date.today(),10)
   self.assertEqual(1,result['sent']);self.assertEqual(1,sms.call_count)
   rows=SmsCampaignRecipient.query.order_by(SmsCampaignRecipient.id).all();self.assertEqual(['발송완료','대기'],[x.status for x in rows])
   self.assertTrue(sms.call_args.args[1].startswith('(광고) 고무신모바일'))

 def test_campaign_failed_delivery_retries_after_backoff(self):
  with app.app_context():
   customer=Customer.query.filter_by(company_code='company-a').first();customer.marketing_consent=True
   campaign=SmsCampaign(company_code='company-a',name='재시도',message='혜택',start_month='2026-09',end_month='2026-09',daily_limit=10,status='진행중',sender_name='고무신모바일',sender_contact='01076671100',opt_out_number='0800000000')
   db.session.add(campaign);db.session.flush();recipient=SmsCampaignRecipient(campaign_id=campaign.id,customer_id=customer.id,phone=customer.phone,scheduled_date=date.today());db.session.add(recipient);db.session.commit();recipient_id=recipient.id
  with app.app_context(),patch('app._send_sms',return_value=False):run_sms_campaigns(date.today(),10)
  with app.app_context():
   recipient=db.session.get(SmsCampaignRecipient,recipient_id);self.assertEqual(('실패',1), (recipient.status,recipient.attempt_count));recipient.next_attempt_at=datetime.utcnow()-timedelta(minutes=1);db.session.commit()
  with app.app_context(),patch('app._send_sms',return_value=True):run_sms_campaigns(date.today(),10)
  with app.app_context():
   recipient=db.session.get(SmsCampaignRecipient,recipient_id);self.assertEqual(('발송완료',2), (recipient.status,recipient.attempt_count))

 def test_campaign_runner_requires_cron_token(self):
  with patch.dict(os.environ,{'CAMPAIGN_CRON_TOKEN':'secret-token'}):
   self.assertEqual(401,self.client.post('/internal/sms-campaigns/run').status_code)
   with patch('app.run_sms_campaigns',return_value={'sent':0,'failed':0}):
    self.assertEqual(200,self.client.post('/internal/sms-campaigns/run',headers={'Authorization':'Bearer secret-token'}).status_code)

 def test_admin_manual_campaign_run_requires_sms_configuration(self):
  self.login_as_a()
  with patch.dict(os.environ,{'SMS_SENDER':'','SMS_OPT_OUT_NUMBER':''}):
   response=self.client.post('/sms-campaigns/run-now',follow_redirects=True)
  self.assertEqual(200,response.status_code);self.assertIn('발송번호와 무료 수신거부번호'.encode(),response.data)

 def test_sale_normalizes_phone_and_auto_transfers_inventory_between_company_branches(self):
  with app.app_context():
   db.session.add(Inventory(serial_number='SERIAL-A',manufacturer='삼성',model='S26',capacity='512GB',color='화이트',branch_id=self.a2_branch,status='보유중'));db.session.commit()
  self.login_as_a()
  response=self.client.post('/sales/new',data={'customer_name':'신규 고객','customer_phone':'010-7667-1100','branch_id':str(self.a_branch),'serial_number':'SERIAL-A','carrier':'LG','opening_type':'번호이동','current_plan':'115'},follow_redirects=False)
  self.assertEqual(302,response.status_code)
  with app.app_context():
   sale=Sale.query.filter_by(serial_number='SERIAL-A').one();inventory=Inventory.query.filter_by(serial_number='SERIAL-A').one();movement=InventoryMovement.query.filter_by(inventory_id=inventory.id).one()
   self.assertEqual('01076671100',sale.customer_phone);self.assertEqual(('S26','512GB','화이트'),(sale.device,sale.storage,sale.color))
   self.assertEqual(self.a_branch,inventory.branch_id);self.assertEqual('판매완료',inventory.status);self.assertEqual('개통자동이관',movement.action)

 def test_sale_cannot_use_another_company_inventory(self):
  with app.app_context():
   db.session.add(Inventory(serial_number='SERIAL-B',manufacturer='삼성',model='S26',branch_id=self.b_branch,status='보유중'));db.session.commit()
  self.login_as_a()
  response=self.client.post('/sales/new',data={'customer_name':'차단 고객','customer_phone':'01099998888','branch_id':str(self.a_branch),'serial_number':'SERIAL-B','carrier':'LG','opening_type':'번호이동'})
  self.assertEqual(403,response.status_code)
  with app.app_context():self.assertEqual('보유중',Inventory.query.filter_by(serial_number='SERIAL-B').one().status)


 def test_pwa_shell_is_installable_without_caching_business_data(self):
  manifest=self.client.get('/static/manifest.webmanifest')
  self.assertEqual(200,manifest.status_code)
  self.assertEqual('standalone',manifest.get_json()['display'])
  worker=self.client.get('/sw.js')
  self.assertEqual(200,worker.status_code)
  body=worker.get_data(as_text=True)
  self.assertIn('/static/offline.html',body)
  self.assertNotIn('/customers',body)
  self.assertNotIn('/sales',body)
  self.login_as_a()
  page=self.client.get('/').get_data(as_text=True)
  self.assertIn('manifest.webmanifest',page)

 def test_admin_readiness_is_admin_only_and_reports_backup(self):
  self.login_as_a()
  response=self.client.get('/admin/readiness')
  self.assertEqual(200,response.status_code)
  body=response.get_data(as_text=True)
  self.assertIn('운영 준비 점검',body)
  self.assertIn('전체 보관백업',body)
  with app.app_context():
   user=db.session.get(User,self.a_user);user.role='staff';db.session.commit()
  self.assertEqual(403,self.client.get('/admin/readiness').status_code)

 def test_admin_dashboard_surfaces_monthly_store_performance(self):
  self.login_as_a();response=self.client.get('/')
  self.assertEqual(200,response.status_code)
  body=response.get_data(as_text=True)
  self.assertIn('매장별 실적',body)
  self.assertIn('직원별 상세 보기',body)
  self.assertIn('이번 달 등록된 매장 실적이 없습니다.',body)

 def test_admin_receives_persistent_backup_warning(self):
  self.login_as_a();response=self.client.get('/notifications')
  body=response.get_data(as_text=True)
  self.assertIn('운영 데이터 보호 확인이 필요합니다.',body)
  self.assertIn('전체 보관백업',body)
  with app.app_context():
   db.session.add(AuditLog(company_code='company-a',username='A관리자',action='관리자 전체백업 다운로드',created_at=datetime.utcnow()));db.session.commit()
  body=self.client.get('/notifications').get_data(as_text=True)
  self.assertNotIn('최근 7일 이내 전체 백업 기록이 없습니다.',body)


if __name__=='__main__':unittest.main()
