import unittest
from datetime import date, timedelta, datetime, timezone
from unittest.mock import patch
import test_manager_permissions as fixtures
import app as module
from app import app, db, Customer, CustomerTask

class HomeWorkTest(unittest.TestCase):
 setUp=fixtures.ManagerPermissionsTest.setUp
 login_as_a=fixtures.ManagerPermissionsTest.login_as_a
 manager=fixtures.ManagerPermissionsTest.manager

 def seed(self):
  with app.app_context():
   own=Customer.query.filter_by(company_code='company-a').first()
   for i in range(8):
    db.session.add(CustomerTask(customer_id=own.id,title='이전약속'+str(i),task_type='요금제 변경',due_date=date(2026,9,26),status='처리예정'))
   db.session.add(CustomerTask(customer_id=own.id,title='오늘요금제예약',task_type='요금제 변경',due_date=date(2026,9,27),status='처리예정'))
   db.session.add(CustomerTask(customer_id=self.b_customer,title='타회사비밀예약',task_type='요금제 변경',due_date=date(2026,9,27),status='처리예정'))
   db.session.commit()

 @patch('app.business_today',return_value=date(2026,9,27))
 def test_today_visible_before_details_even_when_another_date_selected(self,_):
  self.seed();self.login_as_a()
  response=self.client.get('/?date=2026-09-25')
  self.assertEqual(200,response.status_code)
  html=response.data.decode();panel=html[html.index('id="home-today-work"'):html.index('<section class="tm-hero"')]
  self.assertIn('오늘요금제예약',panel);self.assertNotIn('타회사비밀예약',html)
  self.assertIn('고객약속 1건',panel)

 @patch('app.business_today',return_value=date(2026,9,27))
 def test_manager_home_respects_task_permission(self,_):
  self.seed();self.manager({'view':['tasks'],'edit':[],'export':False})
  html=self.client.get('/manager').data.decode()
  self.assertIn('오늘요금제예약',html);self.assertNotIn('타회사비밀예약',html)
  self.assertNotIn('약속 확인 · 처리',html)

 @patch('app.business_today',return_value=date(2026,9,27))
 def test_no_task_permission_no_customer_data(self,_):
  self.seed();self.manager({'view':['inventory'],'edit':[],'export':False})
  html=self.client.get('/manager').data.decode()
  self.assertNotIn('오늘요금제예약',html);self.assertNotIn('home-today-work',html)

 def test_korean_midnight(self):
  with patch('app.datetime') as clock:
   clock.now.side_effect=lambda tz:datetime(2026,9,26,15,1,tzinfo=timezone.utc).astimezone(tz)
   self.assertEqual(date(2026,9,27),module.business_today())

 @patch('app.business_today',return_value=date(2026,9,27))
 def test_excel_reservation_existing_customer_and_reimport(self,_):
  import io
  from openpyxl import Workbook
  self.login_as_a()
  def upload():
   book=Workbook();sheet=book.active
   sheet.append(['고객명','휴대전화','예약일','처리항목'])
   sheet.append(['A 고객','01011112222',datetime(2026,9,27),'요금제 변경'])
   stream=io.BytesIO();book.save(stream);stream.seek(0)
   return self.client.post('/customers/import',data={'branch_id':str(self.a_branch),'file':(stream,'appointments.xlsx')})
  self.assertEqual(302,upload().status_code);self.assertEqual(302,upload().status_code)
  with app.app_context():
   self.assertEqual(1,CustomerTask.query.count())
   self.assertEqual(date(2026,9,27),CustomerTask.query.first().due_date)
  html=self.client.get('/').data.decode()
  self.assertIn('A 고객 · 요금제 변경',html)

 @patch('app.business_today',return_value=date(2026,9,27))
 def test_task_filters_and_pagination_preserve_all_records(self,_):
  self.seed();self.login_as_a()
  with app.app_context():
   own=Customer.query.filter_by(company_code='company-a').first()
   for i in range(205):
    db.session.add(CustomerTask(customer_id=own.id,title=f'페이지약속{i:03}',task_type='고객약속',due_date=date(2026,9,28),status='처리예정'))
   db.session.add(CustomerTask(customer_id=own.id,title='완료확인약속',task_type='고객약속',due_date=date(2026,9,27),status='완료'))
   db.session.commit()
  html=self.client.get('/my-tasks').data.decode()
  self.assertIn('오늘요금제예약',html);self.assertNotIn('이전약속0',html);self.assertNotIn('완료확인약속',html)
  html=self.client.get('/my-tasks?view=upcoming&page=5').data.decode()
  self.assertIn('페이지약속204',html);self.assertIn('205건',html);self.assertNotIn('타회사비밀예약',html)
  html=self.client.get('/my-tasks?view=completed&date=2026-09-27').data.decode()
  self.assertIn('완료확인약속',html);self.assertNotIn('오늘요금제예약',html)
  html=self.client.get('/my-tasks?view=overdue').data.decode()
  self.assertIn('이전약속0',html);self.assertNotIn('오늘요금제예약',html)

 def test_unknown_excel_branch_is_not_silently_reassigned(self):
  import io
  from openpyxl import Workbook
  self.login_as_a();book=Workbook();sheet=book.active
  sheet.append(['고객명','휴대폰번호','예약일','처리항목','처리점'])
  sheet.append(['지점오류고객','01098765432','2026-09-27','요금제','잘못된지점'])
  stream=io.BytesIO();book.save(stream);stream.seek(0)
  response=self.client.post('/customers/import',data={'branch_id':str(self.a_branch),'file':(stream,'appointments.xlsx')})
  self.assertEqual(302,response.status_code)
  with app.app_context():
   self.assertEqual(0,CustomerTask.query.count())
   self.assertIsNone(Customer.query.filter_by(phone='01098765432').first())

 def test_import_explains_missing_date_and_invalid_rows_without_personal_data(self):
  import io
  from openpyxl import Workbook
  self.login_as_a();book=Workbook();sheet=book.active;sheet.title='예약목록'
  sheet.append(['고객명','휴대폰번호','예약일','처리항목'])
  sheet.append(['민감이름','01099887766','','요금제'])
  sheet.append(['날짜오류','01099887765','2026-02-30','요금제'])
  sheet.append([None,None,None,None])
  sheet.append(['정상고객','01099887764','2026-09-28','요금제'])
  stream=io.BytesIO();book.save(stream);stream.seek(0)
  response=self.client.post('/customers/import',data={'branch_id':str(self.a_branch),'file':(stream,'appointments.xlsx')})
  self.assertEqual(302,response.status_code)
  with self.client.session_transaction() as sess:
   messages=' '.join(text for category,text in sess.get('_flashes',[]))
  self.assertIn('예약목록 2행',messages);self.assertIn('예약일이 없음',messages)
  self.assertIn('예약목록 3행',messages);self.assertIn('예약일 오류',messages)
  self.assertIn('제외 2명',messages);self.assertNotIn('민감이름',messages);self.assertNotIn('01099887766',messages)
  with app.app_context():self.assertEqual(1,CustomerTask.query.count())
