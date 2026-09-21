"""Atomic, repeat-safe import of EasySystem inventory and pending customer tasks."""
import hashlib
import io
import json
import re
import secrets
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from html.parser import HTMLParser

from flask import abort, flash, redirect, render_template, request, session, url_for
from sqlalchemy.exc import IntegrityError

INVENTORY_HEADERS = {'기종', '일련번호', '본점입고일', '입고가', '거래처', '서브점'}
TASK_HEADERS = {'고객명', '휴대폰번호', '예약일', '처리항목', '처리점'}
MAX_ROWS = 20000


class ImportProblem(ValueError):
    pass


class HTMLTable(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == 'tr':
            self.row = []
        elif tag in ('td', 'th'):
            self.cell = []
        elif tag == 'br' and self.cell is not None:
            self.cell.append('\n')

    def handle_data(self, value):
        if self.cell is not None:
            self.cell.append(value)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            if self.row is not None:
                self.row.append(''.join(self.cell).strip())
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def json_text(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def fingerprint(row):
    # NO is an export row number and changes between the full and today's lists.
    return hashlib.sha256(json_text({k: v for k, v in row.items() if k != 'NO'}).encode()).hexdigest()


def read_export(filename, content):
    if len(content) > 12 * 1024 * 1024:
        raise ImportProblem('파일 크기는 12MB 이하여야 합니다.')
    tables = []
    if content.startswith(b'PK'):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 128 * 1024 * 1024:
                raise ImportProblem('압축 해제 크기가 너무 큰 엑셀입니다.')
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) > MAX_ROWS:
                        raise ImportProblem('한 파일은 20,000행 이하여야 합니다.')
                    rows.append([cell_text(v) for v in row])
                tables.append(rows)
        finally:
            workbook.close()
    elif content.startswith(b'\xd0\xcf\x11\xe0'):
        raise ImportProblem('바이너리 XLS는 Excel에서 XLSX로 저장한 뒤 올려주세요.')
    else:
        decoded = None
        for encoding in ('utf-8-sig', 'cp949'):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ImportProblem('문자 인코딩을 읽을 수 없습니다. XLSX로 저장해 주세요.')
        parser = HTMLTable()
        parser.feed(decoded)
        tables = [parser.rows]
    outputs = []
    for table in tables:
        nonempty = [row for row in table if any(cell_text(v) for v in row)]
        if not nonempty:
            continue
        header_index = next((i for i, r in enumerate(nonempty[:20])
                             if INVENTORY_HEADERS <= set(r) or TASK_HEADERS <= set(r)), None)
        if header_index is None:
            raise ImportProblem(f'{filename}: 재고 또는 미처리예약 원본의 열 이름을 확인해주세요.')
        header = [cell_text(v) for v in nonempty[header_index]]
        kind = 'inventory' if INVENTORY_HEADERS <= set(header) else 'tasks' if TASK_HEADERS <= set(header) else None
        if not kind:
            raise ImportProblem(f'{filename}: 재고 또는 미처리예약 원본의 열 이름을 확인해주세요.')
        if len(set(header)) != len(header):
            raise ImportProblem(f'{filename}: 중복된 열 이름이 있습니다.')
        records = []
        for number, row in enumerate(nonempty[header_index + 1:], header_index + 2):
            if len(row) != len(header):
                raise ImportProblem(f'{filename} {number}행: 열 개수가 다릅니다. 원본을 확인해주세요.')
            records.append(dict(zip(header, map(cell_text, row))))
        if len(records) > MAX_ROWS:
            raise ImportProblem('한 파일은 20,000행 이하여야 합니다.')
        outputs.append({'filename': filename, 'kind': kind, 'rows': records,
                        'preamble': nonempty[:header_index]})
    if not outputs:
        raise ImportProblem('등록할 행이 없습니다.')
    return outputs


def source_plan(exports):
    stocks = defaultdict(list)
    tasks, max_counts = [], Counter()
    task_source_rows = 0
    source_branches = set()
    for export in exports:
        occurrences = Counter()
        for row in export['rows']:
            branch_key = '서브점' if export['kind'] == 'inventory' else '처리점'
            source_branches.add(row.get(branch_key, ''))
            if export['kind'] == 'inventory':
                serial = row.get('일련번호', '').strip()
                if not serial or not row.get('기종'):
                    raise ImportProblem(f"{export['filename']}: 일련번호 또는 기종이 비어 있습니다.")
                if row not in stocks[serial]:
                    stocks[serial].append(row)
            else:
                task_source_rows += 1
                key = fingerprint(row)
                occurrences[key] += 1
                if occurrences[key] > max_counts[key]:
                    tasks.append({'key': f'{key}:{occurrences[key]}', 'row': row})
        max_counts |= occurrences
    if sum(len(e['rows']) for e in exports) > MAX_ROWS:
        raise ImportProblem('한 번에 20,000행까지만 이전할 수 있습니다.')
    return {'inventory': dict(stocks), 'tasks': tasks, 'branches': sorted(source_branches),
            'conflicts': {k: v for k, v in stocks.items() if len(v) > 1},
            'overlap': task_source_rows - len(tasks), 'source_rows': sum(len(e['rows']) for e in exports)}


def source_date(value, label):
    try:
        return date.fromisoformat(value.replace('/', '-'))
    except (ValueError, AttributeError):
        raise ImportProblem(f'{label}: 날짜를 확인해주세요 ({value or "빈 값"}).')


def source_money(value):
    cleaned = (value or '0').replace(',', '').strip()
    if not re.fullmatch(r'-?\d+', cleaned):
        raise ImportProblem(f'입고가를 확인해주세요: {value}')
    return int(cleaned)


def source_memo(rows):
    return '이지시스템 재고 원본\n' + '\n'.join(json_text(r) for r in rows)


def append_once(existing, addition):
    if addition in (existing or ''):
        return existing
    return ((existing.rstrip() + '\n\n') if existing else '') + addition


def register_easysystem_import(app, db, ns):
    Branch, Customer = ns['Branch'], ns['Customer']
    Inventory, Partner = ns['Inventory'], ns['Partner']
    Task, Movement = ns['CustomerTask'], ns['InventoryMovement']

    class EasySystemImportBatch(db.Model):
        id = db.Column(db.String(64), primary_key=True)
        company_code = db.Column(db.String(50), nullable=False, index=True)
        user_id = db.Column(db.Integer, nullable=False)
        payload = db.Column(db.Text, nullable=False)
        result = db.Column(db.Text)
        created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
        applied_at = db.Column(db.DateTime)

    class EasySystemImportRecord(db.Model):
        id = db.Column(db.Integer, primary_key=True)
        company_code = db.Column(db.String(50), nullable=False, index=True)
        kind = db.Column(db.String(20), nullable=False)
        source_key = db.Column(db.String(120), nullable=False)
        entity_id = db.Column(db.Integer, nullable=False)
        raw_json = db.Column(db.Text, nullable=False)
        __table_args__ = (db.UniqueConstraint('company_code', 'kind', 'source_key', name='uq_easysystem_source'),)

    ns['EasySystemImportBatch'] = EasySystemImportBatch
    ns['EasySystemImportRecord'] = EasySystemImportRecord

    def backup_rows():
        tables = set(db.inspect(db.engine).get_table_names())
        return {key: model.query.filter_by(company_code=ns['current_company']()).all()
                if model.__tablename__ in tables else []
                for key, model in [('easysystem_import_batches', EasySystemImportBatch),
                                   ('easysystem_import_records', EasySystemImportRecord)]}

    ns['easysystem_backup_rows'] = backup_rows

    def context():
        pending = [{'id': r.id, 'row': json.loads(r.raw_json)}
                   for r in EasySystemImportRecord.query.filter_by(company_code=ns['current_company'](), kind='task_pending').all()]
        return {'branches': Branch.query.filter_by(company_code=ns['current_company'](), active=True).order_by(Branch.name).all(),
                'pending': pending, 'csrf': session.setdefault('easysystem_csrf', secrets.token_urlsafe(32))}

    def csrf_check():
        expected = session.get('easysystem_csrf', '')
        if not expected or not secrets.compare_digest(expected, request.form.get('csrf', '')):
            abort(403)

    def batch_owned(token, lock=False):
        query = EasySystemImportBatch.query.filter_by(id=token, company_code=ns['current_company'](), user_id=session['user_id'])
        batch = (query.with_for_update() if lock else query).first_or_404()
        if not batch.applied_at and (datetime.utcnow() - batch.created_at).total_seconds() > 3600:
            raise ImportProblem('미리보기 유효시간이 지났습니다. 파일을 다시 올려주세요.')
        return batch

    def apply_batch(batch, mapping, choices):
        company = ns['current_company']()
        plan = source_plan(json.loads(batch.payload))
        valid_branches = {str(b.id): b for b in Branch.query.filter_by(company_code=company, active=True).all()}
        if any(str(mapping.get(label, '')) not in valid_branches for label in plan['branches']):
            raise ImportProblem('원본의 모든 처리점과 재고 지점을 연결해주세요. 미지정도 보존할 지점을 선택해야 합니다.')
        branch_ids = {b.id for b in valid_branches.values()}
        totals = Counter()
        existing_records = {(r.kind, r.source_key): r for r in EasySystemImportRecord.query.filter_by(company_code=company).all()}
        stocks = {r.serial_number: r for r in Inventory.query.filter(Inventory.branch_id.in_(branch_ids)).all()}
        partners = {p.name: p for p in Partner.query.all()}
        customer_by_key = defaultdict(list)
        for customer in Customer.query.filter_by(company_code=company).all():
            customer_by_key[(re.sub(r'\D', '', customer.phone or ''), customer.branch_id)].append(customer)
        # Validate all records before making any data changes.
        selected = []
        for serial, variants in plan['inventory'].items():
            index = choices.get(serial, 0 if len(variants) == 1 else None)
            if not isinstance(index, int) or not 0 <= index < len(variants):
                raise ImportProblem(f'중복 일련번호 {serial}: 적용할 원본 행을 선택해주세요.')
            row = variants[index]
            bid = valid_branches[str(mapping[row['서브점']])].id
            received = source_date(row['본점입고일'], f'재고 {serial} 입고일')
            cost = source_money(row['입고가'])
            if len(serial) > 100 or len(row['기종']) > 100 or len(row.get('색상', '')) > 50 or len(row.get('거래처', '')) > 100:
                raise ImportProblem(f'재고 {serial}: 입력 길이가 저장 한도를 초과합니다.')
            item = stocks.get(serial)
            if item and (item.sale_id or item.status != '보유중'):
                raise ImportProblem(f'재고 {serial}은 현재 {item.status} 상태입니다. 전체 이전을 중단했습니다.')
            if item and item.branch_id != bid:
                raise ImportProblem(f'재고 {serial}의 현재 지점과 원본 지점이 다릅니다. 이동 여부를 먼저 확인해주세요.')
            selected.append((serial, variants, row, bid, received, cost))
        for entry in plan['tasks']:
            row = entry['row']
            phone = re.sub(r'\D', '', row.get('휴대폰번호', ''))
            if not row.get('고객명') or len(row['고객명']) > 100 or not 10 <= len(phone) <= 15:
                raise ImportProblem('예약 고객명 또는 전화번호를 확인해주세요. 아무 행도 저장하지 않았습니다.')
            if len(row.get('처리자') or row.get('판매자') or '') > 50:
                raise ImportProblem('예약 담당자 이름이 저장 한도를 초과합니다.')

        for serial, variants, row, bid, received, cost in selected:
            partner = None
            if row.get('거래처'):
                partner = partners.get(row['거래처'])
                if partner is None:
                    partner = Partner(name=row['거래처'], category=None, settlement_cycle=None,
                                      memo='이지시스템 재고원본에서 이전. 정산주기 미확인.')
                    db.session.add(partner)
                    db.session.flush()
                    partners[partner.name] = partner
                    totals['partners_created'] += 1
            item = stocks.get(serial)
            created = item is None
            if created:
                item = Inventory(serial_number=serial, branch_id=bid, status='보유중')
                db.session.add(item)
            before = (item.model, item.color, item.received_date, item.purchase_price, item.partner_id, item.memo)
            item.model, item.color = row['기종'], row.get('색상', '')
            item.received_date, item.purchase_price = received, cost
            if partner:
                item.partner_id = partner.id
            item.memo = append_once(item.memo, source_memo(variants))
            after = (item.model, item.color, item.received_date, item.purchase_price, item.partner_id, item.memo)
            db.session.flush()
            stocks[serial] = item
            if created or before != after:
                totals['inventory_created' if created else 'inventory_updated'] += 1
                db.session.add(Movement(inventory_id=item.id, action='이지시스템이전',
                                        from_branch_id=None if created else bid, to_branch_id=bid,
                                        from_status=None if created else item.status, to_status=item.status,
                                        processed_by=session.get('display_name') or session.get('username'),
                                        memo='원본 재고 등록' if created else '원본 거래처·상세기록 보완'))
            else:
                totals['inventory_unchanged'] += 1
            record = existing_records.get(('inventory', serial))
            if record is None:
                record = EasySystemImportRecord(company_code=company, kind='inventory', source_key=serial, entity_id=item.id)
                db.session.add(record)
            record.raw_json = json_text(variants)

        for entry in plan['tasks']:
            key, row = entry['key'], entry['row']
            prior = existing_records.get(('task', key))
            if prior:
                task = db.session.get(Task, prior.entity_id)
                if task is None:
                    raise ImportProblem('이전 기록에 연결된 약속이 삭제되어 있습니다. 기록을 확인해주세요.')
                totals['tasks_existing'] += 1
                continue
            if ('task_pending', key) in existing_records:
                totals['tasks_pending'] += 1
                continue
            bid = valid_branches[str(mapping[row.get('처리점', '')])].id
            phone = re.sub(r'\D', '', row['휴대폰번호'])
            candidates = customer_by_key[(phone, bid)]
            exact = [c for c in candidates if c.name == row['고객명']]
            customer = exact[0] if len(exact) == 1 else None
            if len(exact) > 1:
                raise ImportProblem('같은 전화번호와 지점에 여러 고객이 있습니다. 고객을 정리한 뒤 다시 올려주세요.')
            if customer is None:
                carrier = row.get('통신사', '').upper().replace('SKT', 'SK').replace('LGU+', 'LG')
                customer = Customer(company_code=company, branch_id=bid, name=row['고객명'], phone=phone,
                                    carrier=carrier if carrier in ('SK', 'KT', 'LG') else '',
                                    status='상담중', customer_type='기존손님', marketing_consent=False,
                                    memo='이지시스템 미처리예약 고객. 세부 약속은 고객약속에서 확인.')
                db.session.add(customer)
                db.session.flush()
                customer_by_key[(phone, bid)].append(customer)
                totals['customers_created'] += 1
            try:
                due = source_date(row['예약일'], '예약일')
            except ImportProblem:
                customer.memo = append_once(customer.memo, '날짜 확인 필요 · 이지시스템 예약 원본\n' + json_text(row))
                db.session.add(EasySystemImportRecord(company_code=company, kind='task_pending', source_key=key,
                                                      entity_id=customer.id, raw_json=json_text(row)))
                totals['tasks_pending'] += 1
                continue
            label = row.get('처리항목', '')
            kind = {'요금제': '요금제 변경', '요금제변경': '요금제 변경', '요금제 변경': '요금제 변경',
                    '부가서비스': '부가서비스 해지'}.get(label, label if 0 < len(label) <= 50 else '기타')
            description = '\n'.join(f'{k}: {v}' for k, v in row.items())
            task = Task(customer_id=customer.id, sale_id=None, task_type=kind,
                        title=(row['고객명'] + ' · ' + (label or '고객 약속'))[:150],
                        description='이지시스템 미처리예약 원본\n' + description,
                        due_date=due,
                        assigned_staff=row.get('처리자') or row.get('판매자') or None,
                        status='처리예정', auto_created=False)
            db.session.add(task)
            db.session.flush()
            db.session.add(EasySystemImportRecord(company_code=company, kind='task', source_key=key,
                                                  entity_id=task.id, raw_json=json_text(row)))
            totals['tasks_created'] += 1
        totals.update({'source_rows': plan['source_rows'], 'inventory_total': len(plan['inventory']),
                       'tasks_total': len(plan['tasks']), 'overlap': plan['overlap']})
        batch.result = json_text(dict(totals))
        batch.applied_at = datetime.utcnow()
        # Audit is part of the same transaction. Do not call the existing audit helper,
        # which catches errors by rolling back the caller's transaction.
        db.session.add(ns['AuditLog'](company_code=company, user_id=session['user_id'],
                                      username=session.get('username'), action='이지시스템 일괄이전',
                                      target_type='import', target_id=batch.id, detail=batch.result))
        db.session.commit()
        return dict(totals)

    # Allow tests and operational tooling to verify the exact same transaction path.
    ns['apply_easysystem_batch'] = apply_batch

    @app.route('/admin/easysystem-import', methods=['GET', 'POST'])
    @ns['login_required']
    @ns['admin_required']
    def easysystem_import():
        # Existing signed-in sessions can use the feature immediately after deployment.
        EasySystemImportBatch.__table__.create(db.engine, checkfirst=True)
        EasySystemImportRecord.__table__.create(db.engine, checkfirst=True)
        ctx = context()
        try:
            if request.method == 'POST':
                csrf_check()
                exports = []
                uploads = [f for f in request.files.getlist('files') if f.filename]
                if not 1 <= len(uploads) <= 5:
                    raise ImportProblem('원본 파일을 1~5개 선택해주세요.')
                for upload in uploads:
                    exports.extend(read_export(upload.filename, upload.read()))
                source_plan(exports)
                batch = EasySystemImportBatch(id=secrets.token_urlsafe(24), company_code=ns['current_company'](),
                                              user_id=session['user_id'], payload=json_text(exports))
                db.session.add(batch)
                db.session.commit()
                return redirect(url_for('easysystem_import', batch=batch.id))
            token = request.args.get('batch')
            if token:
                batch = batch_owned(token)
                if batch.applied_at:
                    return render_template('easysystem_import.html', result=json.loads(batch.result), **ctx)
                return render_template('easysystem_import.html', batch=batch, plan=source_plan(json.loads(batch.payload)), **ctx)
        except (ImportProblem, ValueError, zipfile.BadZipFile) as exc:
            db.session.rollback()
            flash(str(exc), 'error')
        return render_template('easysystem_import.html', **ctx)

    @app.post('/admin/easysystem-import/<token>/apply')
    @ns['login_required']
    @ns['admin_required']
    def easysystem_import_apply(token):
        csrf_check()
        try:
            batch = batch_owned(token, lock=True)
            if not batch.applied_at:
                plan = source_plan(json.loads(batch.payload))
                mapping = {label: request.form.get(f'branch_{i}') for i, label in enumerate(plan['branches'])}
                choices = {}
                for i, serial in enumerate(plan['conflicts']):
                    value = request.form.get(f'conflict_{i}', '')
                    if value.isdigit():
                        choices[serial] = int(value)
                apply_batch(batch, mapping, choices)
            return redirect(url_for('easysystem_import', batch=token))
        except (ImportProblem, IntegrityError) as exc:
            db.session.rollback()
            flash(str(exc) if isinstance(exc, ImportProblem) else '다른 자료와의 중복이 발견되어 전체 저장을 취소했습니다. 다시 미리보기를 확인해주세요.', 'error')
            return redirect(url_for('easysystem_import', batch=token))

    @app.post('/admin/easysystem-import/pending/<int:record_id>')
    @ns['login_required']
    @ns['admin_required']
    def easysystem_pending_resolve(record_id):
        csrf_check()
        record = EasySystemImportRecord.query.filter_by(id=record_id, company_code=ns['current_company'](),
                                                        kind='task_pending').with_for_update().first_or_404()
        customer = Customer.query.filter_by(id=record.entity_id, company_code=ns['current_company']()).first_or_404()
        try:
            due = source_date(request.form.get('due_date', ''), '수정 예약일')
            row = json.loads(record.raw_json)
            label = row.get('처리항목') or '고객 약속'
            task = Task(customer_id=customer.id, sale_id=None, task_type=label if len(label) <= 50 else '기타',
                        title=(row['고객명'] + ' · ' + label)[:150],
                        description='이지시스템 예약 원본\n'+'\n'.join(f'{k}: {v}' for k, v in row.items())+
                                    '\n확인된 예약일: '+due.isoformat(),
                        due_date=due, assigned_staff=row.get('처리자') or row.get('판매자') or None,
                        status='처리예정', auto_created=False)
            db.session.add(task)
            db.session.flush()
            record.kind, record.entity_id = 'task', task.id
            db.session.add(ns['AuditLog'](company_code=ns['current_company'](), user_id=session['user_id'],
                                          username=session.get('username'), action='이전 예약일 확인',
                                          target_type='task', target_id=str(task.id), detail='원본 날짜 보존 · 예약일 '+due.isoformat()))
            db.session.commit()
            flash('원본을 보존하고 확인된 날짜로 고객 약속을 등록했습니다.', 'success')
        except (ImportProblem, IntegrityError):
            db.session.rollback()
            flash('날짜 또는 중복 여부를 확인해주세요. 변경 내용을 저장하지 않았습니다.', 'error')
        return redirect(url_for('easysystem_import'))

    return EasySystemImportBatch, EasySystemImportRecord
