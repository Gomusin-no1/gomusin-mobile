"""엑셀 파일 읽기. .xlsx 는 openpyxl, .xls(구형) 는 xlrd 또는 LibreOffice 변환으로 읽는다."""
import os
import shutil
import subprocess
import tempfile

import openpyxl


def load_workbook_any(path):
    """경로의 엑셀 파일을 openpyxl Workbook(값 모드)으로 돌려준다."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return openpyxl.load_workbook(path, data_only=True)
    if ext == ".xls":
        try:
            return _xls_via_xlrd(path)
        except ImportError:
            return _xls_via_libreoffice(path)
    raise ValueError(f"지원하지 않는 파일 형식입니다: {path}")


def _xls_via_xlrd(path):
    """pip install xlrd 가 되어 있으면 변환 없이 바로 읽는다."""
    import xlrd  # noqa: F401  (없으면 ImportError -> LibreOffice 경로로)
    from xlrd import xldate_as_datetime, XL_CELL_DATE

    book = xlrd.open_workbook(path, formatting_info=False)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                val = cell.value
                if cell.ctype == XL_CELL_DATE:
                    try:
                        val = xldate_as_datetime(val, book.datemode)
                    except Exception:
                        pass
                elif isinstance(val, float) and val.is_integer():
                    val = int(val)
                if val == "":
                    val = None
                if val is not None:
                    ws.cell(row=r + 1, column=c + 1, value=val)
        for (r1, r2, c1, c2) in getattr(sh, "merged_cells", []):
            try:
                ws.merge_cells(start_row=r1 + 1, end_row=r2, start_column=c1 + 1, end_column=c2)
            except Exception:
                pass
    return wb


def _xls_via_libreoffice(path):
    """서버에 LibreOffice(soffice)가 설치돼 있으면 .xlsx 로 변환해서 읽는다."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(".xls 파일을 읽으려면 'pip install xlrd' 를 하거나 LibreOffice 를 설치해 주세요.")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "input.xls")
        shutil.copy(path, src)
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", tmp, src],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        return openpyxl.load_workbook(os.path.join(tmp, "input.xlsx"), data_only=True)


def norm_header(v):
    """헤더 비교용: 공백·줄바꿈 제거."""
    return "".join(str(v or "").split())


def clean_label(v):
    """화면 표시용 헤더: 줄바꿈을 공백으로."""
    return " ".join(str(v or "").split())
