"""대사 결과 → 엑셀 보고서 (비교결과 / 미매칭_정산서 / 미매칭_판매일보 / 안내 / 고객별 시트)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .common import fmt_amt
from .vendors import VENDORS

FONT = "Arial"
YELLOW, RED, GREEN, GRAY, HEADFILL = "BF8F00", "C00000", "00B050", "808080", "404040"
FILL = {"match": "FBF2D9", "more": "E2F4E6", "less": "FCE4E4"}
COLOR = {"match": YELLOW, "more": GREEN, "less": RED}
KR = {"match": "일치", "more": "초과(더 받음)", "less": "부족(확인필요)"}
ORDER = {"less": 0, "more": 1, "match": 2}
MONEY = '#,##0;[Red]-#,##0'


def summary_text(it):
    flag = ("  |  ※ 매칭 참고: " + it["flag"]) if it.get("flag") else ""
    if it["status"] == "match":
        return "정산서와 정확히 일치합니다." + flag
    d = "정산서가 더 많이 지급" if it["status"] == "more" else "정산서가 덜 지급(부족)"
    bits = [f"{d} (차액 {fmt_amt(it['diff'])}원)"]
    if it.get("rebate_vendor") is None:
        bits.append(f"정산서 상세: {it['vendor_other_note']} | 판매일보: {it['sale_other_note']}")
    else:
        if it["rebate_diff"]:
            bits.append(f"① 리베이트 차이 {fmt_amt(it['rebate_diff'])}원 (판매일보 {it['rebate_sale']:,.0f} / 정산서 {it['rebate_vendor']:,.0f})")
        else:
            bits.append("① 리베이트는 일치")
        if it["other_diff"]:
            bits.append(f"② 기타조정 차이 {fmt_amt(it['other_diff'])}원 — 정산서: {it['vendor_other_note']} | 판매일보: {it['sale_other_note']}")
    return "  |  ".join(bits) + flag


def _hdr(ws, row, values, widths=None):
    for i, h in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADFILL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_excel(result, out_path):
    sm = result["summary"]
    vendor = VENDORS[sm["vendor_code"]]
    items = result["items"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"비교결과({sm['month']}월)"
    ws["A1"] = f"{sm['vendor']} {sm['year']}년 {sm['month']}월 정산서 vs 판매일보 대사 결과"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = vendor.top_note
    ws["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")

    r = 4
    ws.cell(row=r, column=1, value="■ 총합 비교").font = Font(name=FONT, size=11, bold=True)
    r += 1
    for label, val in [
        (f"판매일보 비교금액(리베이트+구두-차감금액) 합계 — 매칭 {sm['matched']}건", sm["sales_total"]),
        (f"{vendor.totals_vendor_label} — 매칭 {sm['matched']}건", sm["vendor_total"]),
        ("총액 차이 (정산서 − 판매일보)", sm["total_diff"]),
        ("[참고] 리베이트 합계 — 판매일보 R/B", sm["rebate_sales_total"]),
        ("[참고] 리베이트 합계 — 정산서", sm["rebate_vendor_total"]),
    ]:
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10)
        c = ws.cell(row=r, column=7, value=val)
        c.font = Font(name=FONT, size=10, bold=True)
        c.number_format = MONEY
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="■ 건수 요약 (빨간색=부족 건이 맨 위)").font = Font(name=FONT, size=11, bold=True)
    r += 1
    for label, val, color in [("매칭 건수", sm["matched"], None), ("부족(빨간색, 확인필요)", sm["less"], RED),
                              ("초과(초록색, 더 받음)", sm["more"], GREEN), ("일치(노란색)", sm["match"], YELLOW),
                              ("정산서 미매칭 + 매장단위 항목(별도 시트)", sm["unmatched_vendor"] + sm["store_level_items"], None),
                              ("판매일보 미매칭(별도 시트)", sm["unmatched_sales"], None)]:
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10, bold=True)
        c = ws.cell(row=r, column=7, value=val)
        c.font = Font(name=FONT, size=10, bold=True, color=color)
        r += 1
    r += 2

    group_row, header_row = r, r + 1
    for label, c1, c2 in [("비교 결과(종합)", 10, 12), ("① 리베이트 비교", 13, 16),
                          ("② 구두/기타 조정 비교", 17, 21), ("차액 계산에서 제외한 항목 (참고용)", 22, 23)]:
        ws.merge_cells(start_row=group_row, start_column=c1, end_row=group_row, end_column=c2)
        gc = ws.cell(row=group_row, column=c1, value=label)
        gc.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        gc.fill = PatternFill("solid", fgColor="1F1F1F")
        gc.alignment = Alignment(horizontal="center")
    ws.cell(row=group_row, column=1, value="※ 고객명을 누르면 그 고객 판매 건만 따로 보여줍니다 (맞는 항목 초록 / 틀린 항목 빨강)").font = \
        Font(name=FONT, size=9, bold=True, color="0563C1")
    headers = ["No", "상태", "구분", "매장", "판매자", "판매일자", "고객명 (누르면 상세)", "휴대폰번호", "기종",
               "판매일보_비교금액(P+Q-R)", vendor.vendor_cmp_label, "차이(정산서-판매일보)",
               "리베이트_판매일보(P)", vendor.rebate_vendor_label, "리베이트_차이", "리베이트_상태",
               "기타조정_판매일보(Q-R)", vendor.other_vendor_label, "기타조정_차이", "판매일보_상세",
               "정산서_상세(매장 기준 +받음/−차감)", "제외_판매일보", vendor.excl_vendor_label]
    _hdr(ws, header_row, headers, [5, 15, 12, 7, 9, 11, 16, 14, 13, 15, 16, 14, 12, 14, 11, 9, 13, 14, 11, 26, 44, 26, 36])
    money_cols = {10, 11, 12, 13, 14, 15, 17, 18, 19}

    rows_sorted = sorted(items, key=lambda x: (ORDER[x["status"]], x["store"], str(x["date"])))
    # 고객마다 시트를 하나씩 만든다 (고객명을 누르면 그 고객 판매 건만 보이도록)
    cust_sheets, used = [], set()
    for i, it in enumerate(rows_sorted, 1):
        base = re.sub(r"[\[\]:*?/\\']", "", f"{i:03d}_{it['name']}").strip()[:28] or f"{i:03d}"
        nm, k = base, 2
        while nm in used:
            nm, k = f"{base[:26]}_{k}", k + 1
        used.add(nm)
        cust_sheets.append(nm)
    main_rows = []

    rr = header_row + 1
    for i, it in enumerate(rows_sorted, 1):
        has_rb = it["rebate_vendor"] is not None
        dash = lambda v: v if has_rb else "-"  # noqa: E731
        vals = [i, KR[it["status"]], it["kind"], it["store"], it["seller"], it["date"], it["name"], it["phone"], it["device"],
                it["sales_cmp"], it["vendor_cmp"], it["diff"],
                dash(it["rebate_sale"]), dash(it["rebate_vendor"]), dash(it["rebate_diff"]),
                dash("일치" if it["rebate_diff"] == 0 else "확인"),
                dash(it["other_sale"]), dash(it["other_vendor"]), dash(it["other_diff"]),
                it["sale_other_note"], it["vendor_other_note"], it["excluded_sale"], it["excluded_vendor"]]
        fill = PatternFill("solid", fgColor=FILL[it["status"]])
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=rr, column=ci, value=val)
            c.font = Font(name=FONT, size=9, italic=ci in (22, 23), color=GRAY if ci in (22, 23) else None)
            if ci in money_cols:
                c.number_format = MONEY
            if ci <= 2:
                c.fill = fill
        ws.cell(row=rr, column=2).font = Font(name=FONT, size=9, bold=True, color=COLOR[it["status"]])
        c = ws.cell(row=rr, column=12)
        c.font = Font(name=FONT, size=9, bold=True, color=COLOR[it["status"]])
        c.fill = fill
        name_cell = ws.cell(row=rr, column=7)
        name_cell.hyperlink = Hyperlink(ref=name_cell.coordinate, location=f"'{cust_sheets[i - 1]}'!A1")
        name_cell.font = Font(name=FONT, size=9, bold=True, underline="single", color="0563C1")
        if has_rb and it["rebate_diff"]:
            for ci in (15, 16):
                ws.cell(row=rr, column=ci).font = Font(name=FONT, size=9, bold=True, color=RED)
        main_rows.append(rr)
        rr += 1
    ws.freeze_panes = f"D{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(rr - 1, header_row)}"

    ws2 = wb.create_sheet("미매칭_정산서")
    ws2["A1"] = "판매일보에서 대응 항목을 찾지 못한 정산서 라인 + 매장 단위 항목(퀵비 등)"
    ws2["A1"].font = Font(name=FONT, size=11, bold=True)
    _hdr(ws2, 3, vendor.uv_headers, vendor.uv_widths)
    for j, row in enumerate(result["unmatched_vendor"], 4):
        for ci, val in enumerate(row, 1):
            c = ws2.cell(row=j, column=ci, value=val)
            c.font = Font(name=FONT, size=10)
            if isinstance(val, (int, float)) and ci > 2:
                c.number_format = MONEY
    ws2.freeze_panes = "A4"

    ws3 = wb.create_sheet("미매칭_판매일보")
    ws3["A1"] = f"정산서에서 대응 항목을 찾지 못한 판매일보(청약점={vendor.sales_keyword}) 라인"
    ws3["A1"].font = Font(name=FONT, size=11, bold=True)
    _hdr(ws3, 3, ["매장", "판매자", "행", "날짜", "고객명", "기종", "전화번호", "비교금액(P+Q-R)", "정산(W)", "추정 사유"],
         [8, 9, 6, 11, 14, 16, 14, 15, 13, 70])
    for j, s in enumerate(result["unmatched_sales"], 4):
        vals = [s["store"], s["seller"], s["row"], s["date"], s["name"], s["device"], s["phone"], s["sales_cmp"], s["W"], s["reason"]]
        for ci, val in enumerate(vals, 1):
            c = ws3.cell(row=j, column=ci, value=val)
            c.font = Font(name=FONT, size=10)
            if ci in (8, 9):
                c.number_format = MONEY
    ws3.freeze_panes = "A4"

    ws4 = wb.create_sheet("안내")
    ws4.column_dimensions["A"].width = 120
    lines = ["■ 이 파일은 무엇인가요?",
             f"3개 매장 {sm['month']}월 판매일보 중 '청약점={vendor.sales_keyword}' 건과 {sm['vendor']} 정산서를",
             "전화번호 + 고객명(마스킹 패턴) + 날짜 + 기종으로 자동 매칭해 금액을 비교한 결과입니다.",
             "1차로 번호·이름이 모두 맞는 건을 매칭하고, 남은 건은 2차로 이름 오타·번호 오타를 보정해 매칭했습니다.",
             "2차로 매칭된 건은 상세에 '매칭 참고'로 표시되니 한 번 확인해 주세요.",
             "비교결과 시트에서 고객명을 누르면 그 고객 판매 건만 있는 시트(맨 뒤쪽 '001_이름' 탭)로 이동합니다.",
             "그 시트에서는 맞는 항목은 초록 글씨, 틀린 항목은 빨강 글씨로 보이고, 예상 원인과 판매일보 메모가 함께 나옵니다.",
             "'← 목록으로'를 누르면 목록의 그 줄로 돌아가고, '◀ 이전 고객 / 다음 고객 ▶'으로 차례로 볼 수 있습니다. 탭 색: 빨강=부족, 초록=초과, 노랑=일치.",
             ""] + vendor.guide() + [
             "",
             "■ 색상 및 정렬",
             "빨간색 = 정산서가 판매일보보다 적게 줌(부족, 확인 필요) → 맨 위 / 초록색 = 더 줌(초과) / 노란색 = 일치 → 맨 아래",
             "",
             "■ 정산서 정보 (참고)"] + [f"  {x}" for x in sm["vendor_info"]] + [
             "",
             "■ 주의사항",
             "1) 퀵비·유심 구매·환수 같은 매장 단위 항목은 개별 판매 건과 1:1로 맞출 수 없어 '미매칭_정산서' 시트에 따로 둡니다.",
             "2) '미매칭_판매일보'는 이번 정산서에 없는 판매 건입니다. 월말 개통 건은 다음 달 정산서에 들어올 수 있습니다."]
    for i, line in enumerate(lines, 1):
        c = ws4.cell(row=i, column=1, value=line)
        c.font = Font(name=FONT, size=11 if line.startswith("■") else 10, bold=line.startswith("■"))
        c.alignment = Alignment(wrap_text=True)
    for i, (it, nm) in enumerate(zip(rows_sorted, cust_sheets), 1):
        wc = wb.create_sheet(nm)
        wc.sheet_properties.tabColor = {"match": "FFD966", "more": "92D050", "less": "FF7C80"}[it["status"]]
        for col, w in zip("ABCDE", (36, 13, 44, 13, 30)):
            wc.column_dimensions[col].width = w
        nav = [(1, "← 목록으로", f"'{ws.title}'!A{main_rows[i - 1]}")]
        if i > 1:
            nav.append((3, "◀ 이전 고객", f"'{cust_sheets[i - 2]}'!A1"))
        if i < len(cust_sheets):
            nav.append((4, "다음 고객 ▶", f"'{cust_sheets[i]}'!A1"))
        for col, text, loc in nav:
            c = wc.cell(row=1, column=col, value=text)
            c.hyperlink = Hyperlink(ref=c.coordinate, location=loc)
            c.font = Font(name=FONT, size=10, bold=True, underline="single", color="0563C1")
        _detail_block(wc, 3, i, it, ws.title, main_rows[i - 1], vendor)
    wb.save(out_path)
    return out_path


# ------------------------------------------------------------------ 상세 시트 (고객명 클릭 시 이동)
OKC, BADC, INFO, DARK = "00B050", "C00000", "B7791F", "333333"


def _rows_for(it):
    """상세 표의 줄 목록: (판매일보 항목, 판매일보 금액, 정산서 항목, 정산서 금액, 결과, 색)"""
    out = []
    explained = next((c for c in it.get("causes", []) if c["category"] == "정책 단가(리베이트)" and c["amount"] == 0), None)
    if it["rebate_vendor"] is not None:
        if not it["rebate_diff"]:
            out.append(("리베이트(R/B)", it["rebate_sale"], "리베이트", it["rebate_vendor"], "맞음", OKC))
        elif explained:
            out.append(("리베이트(R/B)", it["rebate_sale"], "리베이트", it["rebate_vendor"],
                        f"차이 {fmt_amt(it['rebate_diff'])} — 판매일보 메모로 이미 반영(같은 내용)", OKC))
        else:
            out.append(("리베이트(R/B)", it["rebate_sale"], "리베이트", it["rebate_vendor"], f"다름 ({fmt_amt(it['rebate_diff'])})", BADC))
    mm = it.get("memo_match")
    if mm:
        for memo_lines, labels, amt in mm["pairs"]:
            out.append((" + ".join(memo_lines), amt, " + ".join(labels), amt, "맞음", OKC))
        for v in mm["vendor_only"]:
            out.append(("—", None, v["label"], v["amount"], "정산서에만 있음 (판매일보 메모에 없음)", BADC))
        for m in mm["memo_only"]:
            out.append((m["text"], m["amount"], "—", None, "판매일보에만 있음 (정산서에 없음)", BADC))
    else:
        if it["rebate_vendor"] is not None:
            ok = not it["other_diff"]
            out.append(("구두 − 차감금액 합계", it["other_sale"], "정산서 기타 항목 합계", it["other_vendor"],
                        "맞음" if ok else f"다름 ({fmt_amt(it['other_diff'])})", OKC if ok else BADC))
        else:
            ok = not it["diff"]
            out.append(("판매일보 합계(R/B+구두−차감)", it["sales_cmp"], "정산서 합계", it["vendor_cmp"],
                        "맞음" if ok else f"다름 ({fmt_amt(it['diff'])})", OKC if ok else BADC))
        suspect = {x for c in it.get("causes", []) for x in c.get("items", [])}
        for v in it.get("vendor_items", []):
            if v["category"] == "기본 수수료":
                continue
            bad = v["label"] in suspect
            out.append(("", None, v["label"], v["amount"], "차이 원인으로 보임" if bad else "", BADC if bad else DARK))
    return out


def _detail_block(ws, r, idx, it, main_title, main_row, vendor):
    st_color = COLOR[it["status"]]
    a = ws.cell(row=r, column=1, value=f"{idx}. {it['name']} — {KR[it['status']]}")
    a.font = Font(name=FONT, size=11, bold=True)
    ws.cell(row=r, column=3, value=f"{it['store']} · {it['seller'] or ''} · {it['date']} · {it['device']} · {it['phone'] or ''}").font = \
        Font(name=FONT, size=9, color=GRAY)
    r += 1
    for col, val in ((1, "판매일보 비교금액"), (2, it["sales_cmp"]), (3, "정산서 비교금액"), (4, it["vendor_cmp"]),
                     (5, f"차이 {fmt_amt(it['diff'])}원")):
        c = ws.cell(row=r, column=col, value=val)
        c.font = Font(name=FONT, size=10, bold=True, color=st_color if col == 5 else None)
        if isinstance(val, (int, float)):
            c.number_format = MONEY
    r += 1
    for col, h in enumerate(("판매일보 항목", "금액", "정산서 항목", "금액", "결과"), 1):
        c = ws.cell(row=r, column=col, value=h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADFILL)
    r += 1
    for left, lamt, right, ramt, res, color in _rows_for(it):
        for col, val in enumerate((left, lamt, right, ramt, res), 1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = Font(name=FONT, size=10, color=color, bold=color == BADC)
            if isinstance(val, (int, float)):
                c.number_format = '+#,##0;-#,##0;0'
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    causes = it.get("causes", [])
    if not causes:
        ws.cell(row=r, column=1, value="정산서와 정확히 일치합니다.").font = Font(name=FONT, size=10, bold=True, color=OKC)
        r += 1
    for cz in causes:
        color = INFO if cz["category"] == "매칭" else (OKC if cz["amount"] == 0 else BADC)
        prefix = "매칭 참고" if cz["category"] == "매칭" else ("예상 원인" + ("" if cz["confident"] else "(추정)"))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(row=r, column=1, value=f"{prefix}: {cz['text']}")
        c.font = Font(name=FONT, size=10, color=color, bold=color == BADC)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * max(1, -(-len(c.value) // 95))
        r += 1
    if it.get("sale_memo"):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        memo = "판매일보 메모: " + " / ".join(x.strip() for x in it["sale_memo"].split("\n") if x.strip())
        c = ws.cell(row=r, column=1, value=memo)
        c.font = Font(name=FONT, size=9, italic=True, color=GRAY)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 14 * max(1, -(-len(memo) // 110))
        r += 1
    return r + 1
