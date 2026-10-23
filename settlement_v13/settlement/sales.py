"""매장 판매일보 읽기.

판매일보 양식(3개 매장 공통):
  - 월별 시트: '09월' 또는 '9월'
  - 3행이 헤더, 4행부터 데이터
  - A 날짜(같은 날 두 번째 줄부터는 비어 있음 → 위 행 날짜 사용), B 성명, D 기종, F 청약점(거래처),
    G 개통번호, N 판매자, P R/B(리베이트), Q 구두, R 차감금액, AE 구두내용/부가추가/차감내역(메모),
    S·T 매장마다 이름이 다름(할부프리/추가지원금, 유통망지원금/현금개통), U 유심비, V 수수료, W 정산
  - 정산(W) = P + Q − R − S − T − U − V
"""
import datetime

from openpyxl.utils import column_index_from_string as ci

from .common import digits, num
from .io_excel import load_workbook_any, clean_label

BUNDLE_WORDS = ("인터넷", "TV", "tv", "Tv", "티비", "중고가입", "동판")


def find_month_sheet(wb, month):
    for name in (f"{month:02d}월", f"{month}월"):
        if name in wb.sheetnames:
            return wb[name]
    raise ValueError(f"{month}월 시트를 찾지 못했습니다. 시트 목록: {wb.sheetnames}")


def read_sales(files, month):
    """files: {'본점': 경로, '2호점': 경로, ...}  → 판매일보 행 목록(dict)."""
    rows = []
    for store, path in files.items():
        ws = find_month_sheet(load_workbook_any(path), month)
        s_label = clean_label(ws.cell(3, ci("S")).value) or "S열"
        t_label = clean_label(ws.cell(3, ci("T")).value) or "T열"
        last_date = None
        for r in range(4, ws.max_row + 1):
            g = lambda c: ws.cell(row=r, column=ci(c)).value  # noqa: E731
            if g("A") is not None:
                last_date = g("A")
            if g("B") is None and g("F") is None and g("G") is None:
                continue
            dev = str(g("D") or "")
            rows.append({
                "store": store, "row": r,
                "date": last_date.strftime("%Y-%m-%d") if isinstance(last_date, (datetime.datetime, datetime.date)) else str(last_date),
                "name": str(g("B") or "").strip(),
                "seller": g("N"),
                "device": dev,
                "is_bundle": any(k in dev for k in BUNDLE_WORDS),
                "vendor_raw": str(g("F") or "").replace("\n", " ").strip(),
                "phone": g("G"),
                "phone_digits": digits(g("G")),
                "P": num(g("P")), "Q": num(g("Q")), "R": num(g("R")), "S": num(g("S")), "T": num(g("T")),
                "U": num(g("U")), "V": num(g("V")), "W": num(g("W")),
                "S_label": s_label, "T_label": t_label,
                "memo": str(g("AE") or "").strip(),       # 구두내용/부가추가/차감내역
            })
    return rows
