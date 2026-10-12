"""Input/output boundary for the original v1.3 engine. No replacement calculations."""
import io
import math
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import xlrd

from settlement_v13.settlement.engine import reconcile
from settlement_v13.settlement.sales import read_sales
from settlement_v13.settlement.sales import find_month_sheet
from settlement_v13.settlement.io_excel import load_workbook_any, norm_header
from settlement_v13.settlement.report import write_excel
from settlement_v13.settlement.vendors import VENDORS

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_REQUEST_BYTES = 40 * 1024 * 1024
STORES = ('본점', '2호점', '3호점')


def check_excel(content, filename):
    """Reject wrong formats / oversized workbooks before the original loader opens them."""
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.xlsx', '.xls'):
        raise ValueError('판매일보와 정산서는 .xlsx 또는 .xls 파일을 선택하세요.')
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError('빈 파일은 사용할 수 없으며 파일당 최대 10MB입니다.')
    if suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                members = z.infolist()
                if len(members) > 2000 or sum(i.file_size for i in members) > 50 * 1024 * 1024:
                    raise ValueError('엑셀 압축 해제 크기가 허용 범위를 넘었습니다.')
                if 'xl/workbook.xml' not in z.namelist():
                    raise ValueError('올바른 .xlsx 파일이 아닙니다.')
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                dims = [(ws.max_row or 0, ws.max_column or 0) for ws in wb]
                if len(dims) > 30 or any(r > 20000 or c > 256 for r, c in dims) or sum(r*c for r,c in dims) > 2_000_000:
                    raise ValueError('엑셀 행·열 수가 너무 큽니다. 필요한 월의 데이터만 남긴 사본을 사용하세요.')
            finally:
                wb.close()
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise ValueError('손상되었거나 올바르지 않은 .xlsx 파일입니다.') from exc
    else:
        try:
            wb = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                if wb.nsheets > 30 or any(wb.sheet_by_index(i).nrows > 20000 or wb.sheet_by_index(i).ncols > 256 for i in range(wb.nsheets)):
                    raise ValueError('엑셀 행·열 수가 너무 큽니다.')
            finally:
                wb.release_resources()
        except xlrd.XLRDError as exc:
            raise ValueError('손상되었거나 올바르지 않은 .xls 파일입니다.') from exc
    return suffix


def validate_sales_rows(rows, year, month, vendor_codes):
    relevant = [s for s in rows if any(VENDORS[c].sales_keyword in s['vendor_raw'] for c in vendor_codes)]
    if len(relevant) > 2500:
        raise ValueError('한 번에 대조할 판매는 최대 2,500건입니다. 파일을 나눠 주세요.')
    for s in relevant:
        try:
            d = date.fromisoformat(s['date'])
        except (ValueError, TypeError) as exc:
            raise ValueError(f"{s['store']} 판매일보 {s['row']}행: 날짜를 확인하세요.") from exc
        if (d.year, d.month) != (year, month):
            raise ValueError(f"{s['store']} 판매일보 {s['row']}행의 날짜가 선택한 정산월과 다릅니다.")
    return rows


def calculate(year, month, sales_files, vendor_files):
    """Return original result dicts in original adapter order, without changing amounts or matches."""
    for store, path in sales_files.items():
        wb = load_workbook_any(path)
        try:
            ws = find_month_sheet(wb, month)
            if norm_header(ws['F3'].value) != '청약점' or norm_header(ws['P3'].value).upper() not in ('R/B', 'RB', '리베이트'):
                raise ValueError(f'{store}: 판매일보 양식의 F3 청약점 / P3 R/B 열을 확인하세요.')
            formula_wb = openpyxl.load_workbook(path, read_only=True, data_only=False) if Path(path).suffix == '.xlsx' else None
            try:
                formulas = find_month_sheet(formula_wb, month) if formula_wb else None
                # Read formula cells once; repeated read-only cell access is quadratic.
                formula_rows = {i: row for i, row in enumerate(formulas.iter_rows(min_row=4, min_col=16, max_col=18), 4)} if formulas else {}
                for row in range(4, ws.max_row + 1):
                    vendor = str(ws.cell(row, 6).value or '')
                    if not any(VENDORS[c].sales_keyword in vendor for c in vendor_files):
                        continue
                    for col in (16, 17, 18):
                        cell = ws.cell(row, col)
                        v = cell.value
                        f = formula_rows.get(row, (None, None, None))[col-16]
                        if v is None and f is not None and f.data_type == 'f':
                            raise ValueError(f'{store} {cell.coordinate}: 수식 계산값이 없습니다. 엑셀에서 재계산·저장한 뒤 업로드하세요.')
                        if v not in (None, '') and (isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or cell.data_type == 'e'):
                            raise ValueError(f'{store} {cell.coordinate}: 금액을 숫자로 입력한 뒤 저장하세요.')
            finally:
                if formula_wb:
                    formula_wb.close()
        finally:
            wb.close()
    rows = validate_sales_rows(read_sales(sales_files, month), year, month, vendor_files)
    results = []
    for code, paths in vendor_files.items():
        try:
            result = reconcile(code, year, month, sales_files, paths, sales_rows=rows)
        except (KeyError, StopIteration, IndexError) as exc:
            raise ValueError(f'{VENDORS[code].name} 정산서의 필수 시트·열 제목을 확인하세요. 선택한 거래처 양식과 맞지 않습니다.') from exc
        results.append(result)
    return results


def screen_data(results):
    return [dict(r, label=r['summary']['vendor'], month=f"{r['summary']['month']}월",
                 uv_headers=VENDORS[r['summary']['vendor_code']].uv_headers) for r in results]


def excel_report(result):
    """Keep original workbook layout; uploaded text must not become executable formulas."""
    out = io.BytesIO()
    write_excel(result, out)
    out.seek(0)
    wb = openpyxl.load_workbook(out)
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    safe = io.BytesIO()
    wb.save(safe)
    wb.close()
    safe.seek(0)
    return safe
