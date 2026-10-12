"""Settlement import, conservative matching and explicit watch profit arithmetic.

All amounts are whole KRW. No workbook formula is executed on the server.
"""
import calendar
import io
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from zoneinfo import ZoneInfo

MAX_ROWS = 10000
MAX_COLS = 200
BASIS = {'gross': '판매일보 세전 정산금', 'net': '판매일보 세후 정산금',
         'commission': '리베이트 + 구두추가 − 차감금액'}
ALIASES = {
    'name': ['고객명', '가입자명', '명의자', '가입자'],
    'phone': ['개통번호', 'ctn', '휴대폰번호', '휴대전화', '전화번호', '고객번호'],
    'date': ['개통시간', '판매일자', '개통일', '개통일자', '판매일'],
    'serial': ['일련번호', '일련', 'imei', '단말일련번호'],
    'amount': ['정산금액', '실지급계', '판매점기준실제이윤', '실지급액', '지급계', '무선수수료'],
}


def today_kst():
    return datetime.now(ZoneInfo('Asia/Seoul')).date()


def norm(value):
    return re.sub(r'[^a-z0-9가-힣]', '', unicodedata.normalize('NFKC', str(value or '')).lower())


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return '' if value is None else str(value)


def won(value, *, allow_negative=True):
    raw = str(value if value is not None else '').strip().replace(',', '').replace('원', '').replace('₩', '').strip()
    if raw.startswith('(') and raw.endswith(')'):
        raw = '-' + raw[1:-1]
    if not re.fullmatch(r'[+-]?\d+(?:\.0+)?', raw):
        raise ValueError('금액은 원 단위 정수로 입력하세요. 빈칸은 0원으로 간주하지 않습니다.')
    amount = int(Decimal(raw))
    if abs(amount) > 100_000_000_000 or (not allow_negative and amount < 0):
        raise ValueError('금액 범위를 확인하세요.')
    return amount


def phone(value):
    raw = scalar(value).strip()
    if any(c in raw for c in '*xX●•'):
        return ''
    digits = re.sub(r'\D', '', raw)
    if len(digits) == 10 and digits.startswith('10'):
        digits = '0' + digits
    return digits if len(digits) in (10, 11) and digits.startswith('0') else ''


def serial(value):
    raw = scalar(value).strip()
    return '' if not raw or any(c in raw for c in '*●•') else norm(raw)


def parsed_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or '').strip()
    match = re.match(r'^(\d{4})[-./](\d{1,2})[-./](\d{1,2})(?:\D|$)', raw)
    if not match:
        match = re.fullmatch(r'(\d{4})(\d{2})(\d{2})', raw)
    if match:
        try:
            return date(*map(int, match.groups()))
        except ValueError:
            pass
    return None


def read_tables(content, filename):
    """Return bounded cells from .xls/.xlsx with dates normalized, no macros/eval."""
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('파일은 10MB 이하만 업로드할 수 있습니다.')
    suffix = filename.lower().rsplit('.', 1)[-1]
    try:
        tables = []
        if suffix == 'xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('압축 해제 크기가 너무 큽니다. 시트를 나누어 업로드하세요.')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            if len(book.worksheets) > 30:
                raise ValueError('시트 수는 30개 이하로 나누어 주세요.')
            for sheet in book:
                if (sheet.max_row or 0) > MAX_ROWS + 50 or (sheet.max_column or 0) > MAX_COLS:
                    raise ValueError('시트는 10,000행·200열 이하로 나누어 주세요.')
                rows = []
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= MAX_ROWS + 50 or len(row) > MAX_COLS:
                        raise ValueError('시트 크기 제한을 초과했습니다.')
                    rows.append([scalar(v) for v in row])
                tables.append({'name': sheet.title, 'rows': rows})
            book.close()
        elif suffix == 'xls':
            import xlrd
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            if book.nsheets > 30:
                raise ValueError('시트 수는 30개 이하로 나누어 주세요.')
            for sheet in book.sheets():
                if sheet.nrows > MAX_ROWS + 50 or sheet.ncols > MAX_COLS:
                    raise ValueError('시트는 10,000행·200열 이하로 나누어 주세요.')
                rows = []
                for i in range(sheet.nrows):
                    row = []
                    for cell in sheet.row(i):
                        value = xlrd.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                        row.append('엑셀 오류' if cell.ctype == xlrd.XL_CELL_ERROR else scalar(value))
                    rows.append(row)
                tables.append({'name': sheet.name, 'rows': rows})
            book.release_resources()
        else:
            raise ValueError('지원 형식은 .xls와 .xlsx입니다.')
        return tables
    except (ValueError, ImportError):
        raise
    except Exception as exc:
        raise ValueError('엑셀을 읽을 수 없습니다. 암호를 해제하고 정상 .xls 또는 .xlsx 파일로 저장하세요.') from exc


def detect_mapping(headers):
    clean = [norm(x) for x in headers]
    result = {}
    for field, aliases in ALIASES.items():
        matches = [i for i, v in enumerate(clean) if v and v in aliases]
        # Multiple amounts require an explicit selection, never a guessed cash basis.
        result[field] = matches[0] if len(matches) == 1 else None
    return result


def header_candidates(tables):
    found = []
    for si, table in enumerate(tables):
        for ri, row in enumerate(table['rows'][:50]):
            mapping = detect_mapping(row)
            if mapping['date'] is not None and (mapping['phone'] is not None or mapping['serial'] is not None):
                found.append({'sheet': si, 'header': ri, 'name': table['name'], 'headers': row,
                              'mapping': mapping, 'data_count': max(0, len(table['rows']) - ri - 1)})
                break
    return found


def statement_rows(tables, config, month):
    si, hi = int(config['sheet']), int(config['header'])
    if not 0 <= si < len(tables) or not 0 <= hi < min(50, len(tables[si]['rows'])):
        raise ValueError('시트와 제목 행을 확인하세요.')
    table = tables[si]
    mapping = config['mapping']
    width = len(table['rows'][hi])
    for key in ['date', 'amount']:
        if mapping.get(key) is None:
            raise ValueError('개통일과 정산금액 열은 필수입니다.')
    if mapping.get('phone') is None and mapping.get('serial') is None:
        raise ValueError('전화번호 또는 일련번호 열을 선택하세요.')
    indices = [int(v) for v in mapping.values() if v is not None]
    if len(indices) != len(set(indices)) or any(v < 0 or v >= width for v in indices):
        raise ValueError('열 선택이 중복되거나 범위를 벗어났습니다.')
    result = []
    for row_no, cells in enumerate(table['rows'][hi + 1:], hi + 2):
        if not any(str(c).strip() for c in cells):
            continue
        def get(key):
            index = mapping.get(key)
            return cells[int(index)] if index is not None and int(index) < len(cells) else ''
        # Summary-only lines cannot be mistaken for transactions.
        name, raw_phone, raw_serial, raw_date = get('name'), get('phone'), get('serial'), get('date')
        if not raw_date and not raw_phone and not raw_serial and (not name or any(x in norm(name) for x in ['합계', '소계', '총계'])):
            continue
        d = parsed_date(raw_date)
        error = ''
        try:
            amount = won(get('amount')) * int(config.get('sign', 1))
        except ValueError:
            amount = None
            error = '금액 빈칸·오류 또는 수식 계산값 없음'
        if not d:
            error = '개통일 확인 필요'
        elif d.strftime('%Y-%m') != month:
            error = '대조월 외 거래 (별도 월로 대조 필요)'
        if not phone(raw_phone) and not serial(raw_serial):
            error = '전체 전화번호·일련번호 확인 필요'
        result.append({'row': row_no, 'sheet': table['name'], 'name': str(name)[:100],
                       'phone': phone(raw_phone), 'serial': serial(raw_serial),
                       'date': d.isoformat() if d else str(raw_date)[:30], 'actual': amount,
                       'error': error})
    return result


def reconcile(rows, sales):
    """Unique exact date+identifier only; reserve all ambiguous candidates from missing."""
    output, involved = [], set()
    by_phone, by_serial = defaultdict(dict), defaultdict(dict)
    for sale in sales:
        for p in sale['phones']:
            by_phone[(sale['date'], p)][sale['id']] = sale
        if sale['serial']:
            by_serial[(sale['date'], sale['serial'])][sale['id']] = sale
    for row in rows:
        item = dict(row, sale_id=None, expected=None, delta=None)
        candidates = []
        if not row['error']:
            ids = dict(by_phone.get((row['date'], row['phone']), {})) if row['phone'] else {}
            if row['serial']:
                ids.update(by_serial.get((row['date'], row['serial']), {}))
            candidates = list(ids.values())
            involved.update(s['id'] for s in candidates)
        item['candidate_ids'] = [s['id'] for s in candidates]
        if row['error']:
            item['status'] = '확인필요'
        elif not candidates:
            item['status'] = '판매미등록'
        elif len(candidates) > 1:
            item['status'] = '중복후보'
        else:
            sale = candidates[0]
            mismatch = (row['name'] and sale['name'] and norm(row['name']) != norm(sale['name'])) or (
                row['phone'] and sale['phones'] and row['phone'] not in sale['phones']) or (
                row['serial'] and sale['serial'] and row['serial'] != sale['serial'])
            if sale.get('expected') is None or sale.get('sale_status') in ('취소', '개통취소', '반품', '해지'):
                item['status'] = '확인필요'
                item['error'] = '판매일보 기준금액 또는 취소·반품 상태 확인'
            elif mismatch:
                item['status'] = '정보불일치'
                item['error'] = '이름·전화번호·일련번호 중 불일치 항목 확인'
            else:
                item.update(sale_id=sale['id'], expected=sale['expected'])
                item['delta'] = row['actual'] - sale['expected']
                item['status'] = '정상' if item['delta'] == 0 else ('부족' if item['delta'] < 0 else '초과')
        output.append(item)
    counts = Counter(x['sale_id'] for x in output if x['sale_id'] is not None)
    for item in output:
        if item['sale_id'] is not None and counts[item['sale_id']] > 1:
            item.update(status='중복행', expected=None, delta=None, sale_id=None, error='동일 판매에 여러 정산 행이 있습니다. 자동 합산하지 않았습니다.')
    for sale in sales:
        if sale['id'] not in involved:
            output.append({'row': None, 'sheet': '', 'name': sale['name'], 'phone': next(iter(sale['phones']), ''),
                           'serial': sale['serial'], 'date': sale['date'], 'actual': None,
                           'expected': sale['expected'], 'delta': None, 'sale_id': sale['id'],
                           'candidate_ids': [], 'status': '정산서누락', 'error': '정산서에서 찾지 못함. 미입금 확정이 아닙니다.'})
    return output


def next_month_end(value):
    year, month = (value.year + 1, 1) if value.month == 12 else (value.year, value.month + 1)
    return date(year, month, calendar.monthrange(year, month)[1])


def watch_calculate(values):
    def amount(key):
        return won(values[key], allow_negative=False)
    rebate, device = amount('rebate'), amount('device_cost')
    sold = None if values.get('resale_price') in [None, ''] else amount('resale_price')
    penalty, customer, fee = amount('penalty'), amount('customer_payment'), amount('seller_fee')
    if any(not re.fullmatch(r'\d{1,3}(?:\.\d{1,2})?', str(values.get(k, d))) for k, d in [('tax_rate','13.3'),('share_rate','50')]):
        raise ValueError('비율은 소수 둘째 자리까지 숫자로 입력하세요.')
    try:
        rate = Decimal(str(values.get('tax_rate', '13.3')))
        share = Decimal(str(values.get('share_rate', '50')))
    except InvalidOperation as exc:
        raise ValueError('비율을 숫자로 입력하세요.') from exc
    if not rate.is_finite() or not share.is_finite() or not 0 <= rate <= 100 or not 0 <= share <= 100:
        raise ValueError('비율은 0~100 사이여야 합니다.')
    mode = values.get('penalty_tax_mode', 'deduct')
    if mode not in ('deduct', 'full'):
        raise ValueError('위약금 공제 방식을 확인하세요.')
    rounding = lambda n: int(n.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    base = rebate - device
    tax = rounding(Decimal(max(base, 0)) * rate / 100)
    penalty_credit = rounding(Decimal(penalty) * rate / 100) if mode == 'deduct' else 0
    expense = penalty - penalty_credit
    profit = None if sold is None else base - tax + sold - expense - customer - fee
    staff_share = None if profit is None else int(Decimal(max(profit, 0)) * share / 100)
    return {'base': base, 'tax': tax, 'net_rebate': base - tax, 'penalty_credit': penalty_credit,
            'penalty_expense': expense, 'profit': profit, 'staff_share': staff_share,
            'staff_total': None if profit is None else fee + staff_share,
            'store_share': None if profit is None else profit - staff_share}
